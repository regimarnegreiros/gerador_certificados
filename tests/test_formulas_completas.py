"""
Testes específicos para validação das fórmulas dinâmicas e formatação condicional
em todas as abas da planilha do relatório de presenças (modelo híbrido Python + Excel).
"""

import unittest
from pathlib import Path
import tempfile
import pandas as pd
import openpyxl

from core.presenca import (
    criar_aba_presencas,
    criar_aba_informacoes_curso,
    criar_aba_interacoes_participante,
    criar_aba_interacoes_encontro,
    criar_aba_indicadores,
    criar_aba_custos,
    organizar_colunas,
    exportar_excel,
)
from core.certificados import carregar_relatorio_presenca, filtrar_aptos


class TestFormulasCompletas(unittest.TestCase):

    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        self.saida_excel = self.temp_dir / "Relatorio_Formulas_Teste.xlsx"

        self.config = {
            "NOME_DO_CURSO": "Curso de Teste Fórmulas",
            "CARGA_HORARIA": 10,
            "CUSTO_TOTAL_EM_REAIS": 2000.00,
            "PERCENTUAL_MINIMO_CARGA": 0.75,
            "PERCENTUAL_MINIMO_ENCONTROS": 0.75,
            "TEMPO_MINIMO_PRESENCA_MINUTOS": 20,
        }

        # 3 participantes:
        # 1. Ana: 2 encontros, 8h total -> APTO
        # 2. Beto: 1 encontro, 4h total -> NÃO APTO
        # 3. Caio: 0 encontros, 0h total -> NÃO APTO
        self.relatorio = pd.DataFrame({
            "Nome": ["Ana Silva", "Beto Souza", "Caio Lima"],
            "UORG": ["DIR", "SGP", "DITEC"],
            "Email": ["ana@teste.gov.br", "beto@teste.gov.br", "caio@teste.gov.br"],
            "Encontro 1": ["04:00:00", "04:00:00", "00:00:00"],
            "Comparecimento 1": ["Presente", "Presente", "Ausente"],
            "Encontro 2": ["04:00:00", "00:00:00", "00:00:00"],
            "Comparecimento 2": ["Presente", "Ausente", "Ausente"],
            "Presenças": [2, 1, 0],
            "CH Total": ["8:00:00", "4:00:00", "0:00:00"],
            "% da CH Total": [0.8, 0.4, 0.0],
            "Apto por CH": ["APTO", "NÃO APTO", "NÃO APTO"],
            "Apto por Presença": ["APTO", "NÃO APTO", "NÃO APTO"],
            "Apto em Ambos": ["APTO", "NÃO APTO", "NÃO APTO"],
        })

        self.df_presencas = criar_aba_presencas(self.relatorio, quantidade_encontros=2)
        self.df_certificacao = organizar_colunas(self.relatorio, quantidade_encontros=2)

        # Dados sintéticos de interações
        df_int_e1 = pd.DataFrame({
            "Nome": ["Ana Silva", "Beto Souza", "Caio Lima"],
            "camera_ligou": [1, 0, 0],
            "microfone_ligou": [1, 0, 0],
            "maos_qtd": [2, 0, 0],
            "reacoes_qtd": [5, 0, 0],
        })
        df_int_e2 = pd.DataFrame({
            "Nome": ["Ana Silva", "Beto Souza", "Caio Lima"],
            "camera_ligou": [1, 0, 0],
            "microfone_ligou": [0, 0, 0],
            "maos_qtd": [1, 0, 0],
            "reacoes_qtd": [3, 0, 0],
        })
        df_res_e1 = pd.DataFrame({
            "Nome": ["Ana Silva", "Beto Souza", "Caio Lima"],
            "Compareceu": ["Presente", "Presente", "Ausente"],
            "Duração da Reunião": ["04:00:00", "04:00:00", "00:00:00"],
        })
        df_res_e2 = pd.DataFrame({
            "Nome": ["Ana Silva", "Beto Souza", "Caio Lima"],
            "Compareceu": ["Presente", "Ausente", "Ausente"],
            "Duração da Reunião": ["04:00:00", "00:00:00", "00:00:00"],
        })

        lista_int = [df_int_e1, df_int_e2]
        lista_res = [df_res_e1, df_res_e2]
        lista_resumos = [
            {"data": "10/09/2026", "duracao_segundos": 14400, "duracao_str": "04:00:00", "tempo_medio_segundos": 14400, "tempo_medio_str": "04:00:00"},
            {"data": "11/09/2026", "duracao_segundos": 14400, "duracao_str": "04:00:00", "tempo_medio_segundos": 14400, "tempo_medio_str": "04:00:00"},
        ]

        self.df_part = criar_aba_interacoes_participante(self.relatorio, lista_int, lista_res)
        self.df_enc = criar_aba_interacoes_encontro(3, lista_resumos, lista_int, lista_res, config=self.config)
        self.df_ind = criar_aba_indicadores(3, self.relatorio, self.df_part, self.df_enc, lista_resumos, lista_res)
        self.df_custos = criar_aba_custos(3, self.relatorio, 2, 10.0, lista_resumos, lista_res, config=self.config)
        self.df_info = criar_aba_informacoes_curso(self.config)

        exportar_excel(
            relatorio=self.df_certificacao,
            caminho_saida=self.saida_excel,
            df_presencas=self.df_presencas,
            df_informacoes=self.df_info,
            df_interacoes_participante=self.df_part,
            df_interacoes_encontro=self.df_enc,
            df_indicadores=self.df_ind,
            df_custos=self.df_custos,
            datas_encontros=["10/09/2026", "11/09/2026"],
        )

        self.wb = openpyxl.load_workbook(self.saida_excel)

    def test_formulas_aba_certificacao(self):
        ws = self.wb["Certificação"]

        # Linha 2 (Ana Silva)
        # Presenças: COUNTIFS dinâmico no intervalo B2:E2
        formula_pres = str(ws["D2"].value)
        self.assertEqual(formula_pres, '=COUNTIFS(\'Presenças\'!B2:E2, "?*", \'Presenças\'!B2:E2, "<>Ausente")')

        # CH Total: soma de durações via SUMIF no cabeçalho B1:E1
        self.assertEqual(ws["E2"].value, "=SUMIF('Presenças'!$B$1:$E$1, \"Encontro*\", 'Presenças'!B2:E2)")
        self.assertEqual(ws["E2"].number_format, "[h]:mm:ss")

        # % da CH Total: (E2*24) / Informações do Curso!$B$3
        self.assertEqual(ws["F2"].value, "=IF('Informações do Curso'!$B$3>0, (E2*24)/'Informações do Curso'!$B$3, 0)")
        self.assertEqual(ws["F2"].number_format, "0.0%")

        # Apto por CH: IF(F2 >= Informações do Curso!$B$8)
        self.assertEqual(ws["G2"].value, '=IF(F2>=\'Informações do Curso\'!$B$8, "APTO", "NÃO APTO")')

        # Apto por Presença: IF(D2 >= MAX(1, ROUNDUP(Custos!$B$6 * Informações do Curso!$B$9, 0)))
        formula_apto_pres = str(ws["H2"].value)
        self.assertTrue(formula_apto_pres.startswith("=IF(D2>=MAX(1, ROUNDUP('Custos'!$B$6 * 'Informações do Curso'!$B$9, 0)), \"APTO\", \"NÃO APTO\")"))

        # Apto em Ambos: IF(AND(G2="APTO", H2="APTO"), "APTO", "NÃO APTO")
        self.assertEqual(ws["I2"].value, '=IF(AND(G2="APTO", H2="APTO"), "APTO", "NÃO APTO")')

    def test_formatacao_condicional_certificacao(self):
        ws = self.wb["Certificação"]
        # Verifica se foram registradas regras de formatação condicional cobrindo G2:I4
        cf_regras = list(ws.conditional_formatting)
        self.assertGreater(len(cf_regras), 0)

        faixas = [str(r.sqref) for r in cf_regras]
        self.assertTrue(any("G2:I" in f for f in faixas))

        # Verifica se as regras possuem preenchimento (fill) com start_color e end_color e fonte (font)
        regras = [r for cf in cf_regras for r in cf.rules]
        self.assertGreaterEqual(len(regras), 2)
        for r in regras:
            self.assertIsNotNone(r.dxf.fill)
            self.assertEqual(r.dxf.fill.fill_type, "solid")
            self.assertIsNotNone(r.dxf.fill.start_color)
            self.assertIsNotNone(r.dxf.fill.end_color)
            self.assertIsNotNone(r.dxf.font)
            self.assertIsNotNone(r.dxf.font.color)

    def test_formulas_aba_interacoes_participante(self):
        ws = self.wb["Interações por participante"]

        # Qtd Presenças vinculada à Certificação
        self.assertEqual(ws["B2"].value, "='Certificação'!D2")

        # % Câmera: IF(B2>0, C2/B2, 0)
        self.assertEqual(ws["D2"].value, "=IF(B2>0, C2/B2, 0)")
        self.assertEqual(ws["D2"].number_format, "0.0%")

        # % Microfone: IF(B2>0, E2/B2, 0)
        self.assertEqual(ws["F2"].value, "=IF(B2>0, E2/B2, 0)")
        self.assertEqual(ws["F2"].number_format, "0.0%")

    def test_formulas_aba_interacoes_encontro(self):
        ws = self.wb["Interações por encontro"]

        # Encontro 1 (linha 2): Tempo Médio e Presentes
        self.assertEqual(ws["D2"].value, '=IFERROR(AVERAGEIF(\'Presenças\'!B:B, ">0"), 0)')
        self.assertEqual(ws["D2"].number_format, "hh:mm:ss")

        formula_pres = str(ws["E2"].value)
        self.assertTrue(formula_pres.startswith("=COUNTIFS('Presenças'!C:C, \"<>Ausente\""))
        self.assertIn("de", formula_pres)
        self.assertIn("Indicadores!$B$2", formula_pres)

        # Colunas de câmera, microfone, mãos e reações SEM fórmulas (pré-calculadas)
        self.assertFalse(str(ws["F2"].value).startswith("="))
        self.assertFalse(str(ws["G2"].value).startswith("="))
        self.assertFalse(str(ws["H2"].value).startswith("="))
        self.assertFalse(str(ws["I2"].value).startswith("="))

    def test_formulas_indicadores_e_custos(self):
        ws_ind = self.wb["Indicadores"]
        mapa_ind = {str(ws_ind.cell(row=r, column=1).value): ws_ind.cell(row=r, column=2).value for r in range(2, ws_ind.max_row + 1)}

        # Tempo médio de participação com fórmula AVERAGEIF em Certificação
        self.assertEqual(mapa_ind["Tempo médio de participação"], '=IFERROR(AVERAGEIF(\'Certificação\'!E:E, ">0"), 0)')
        # Carga horária realizada sem fórmula
        self.assertEqual(mapa_ind["Carga horária realizada"], "08:00:00")

        ws_custos = self.wb["Custos"]
        mapa_custos = {str(ws_custos.cell(row=r, column=1).value): ws_custos.cell(row=r, column=2).value for r in range(2, ws_custos.max_row + 1)}

        # Carga horária vinculada a Informações do Curso
        self.assertEqual(mapa_custos["Carga horária"], "='Informações do Curso'!$B$3/24")
        # Tempo médio de participação realizado vinculado a Indicadores
        self.assertEqual(mapa_custos["Tempo médio de participação realizado"], "=Indicadores!B6")
        # Carga horária realizada sem fórmula
        self.assertEqual(mapa_custos["Carga horária realizada"], "08:00:00")

    def test_leitura_hibrida_sem_excel(self):
        # Valida que o pipeline carrega o arquivo recém-gerado sem abrir no Excel
        # e reconhece perfeitamente os participantes aptos graças ao cache XML e fallbacks
        df_rel = carregar_relatorio_presenca(self.saida_excel)
        aptos, nao_aptos = filtrar_aptos(df_rel)

        self.assertEqual(len(aptos), 1)
        self.assertEqual(aptos.iloc[0]["Nome"], "Ana Silva")
        self.assertEqual(len(nao_aptos), 2)


if __name__ == "__main__":
    unittest.main()

