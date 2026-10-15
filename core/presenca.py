"""
Módulo para processamento de listas de presença do Microsoft Teams,
cruzamento com dados de servidores ANM e geração do relatório Excel formatado.
Extraído e adaptado diretamente de presensa_teams.ipynb.
"""

from pathlib import Path
import unicodedata
import os
import re
import math
import csv
import io
from datetime import datetime, timedelta
import pandas as pd
from openpyxl.styles import (
    PatternFill,
    Font,
    Alignment,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from typing import Tuple, List, Dict, Any, Optional, Set, Collection
from core.caminhos import (
    caminho_estendido,
    caminho_estendido_path,
    garantir_diretorio,
    abrir_arquivo,
    copiar_arquivo_com_fallback,
    existe_com_fallback,
    listar_diretorio_com_fallback,
    ler_excel_com_fallback,
    sanitizar_nome_arquivo,
    encurtar_caminho_se_necessario,
)


def tempo_em_segundos(tempo: Any) -> int:
    """Converte strings de tempo (HH:MM:SS ou 1h 30m 20s) em segundos."""
    if pd.isna(tempo):
        return 0

    tempo = str(tempo).strip()

    # Formato HH:MM:SS ou MM:SS
    if ":" in tempo:
        partes = tempo.split(":")
        if len(partes) == 3:
            horas, minutos, segundos = map(int, partes)
            return horas * 3600 + minutos * 60 + segundos
        elif len(partes) == 2:
            minutos, segundos = map(int, partes)
            return minutos * 60 + segundos

    # Formato 1h 30m 20s
    valores = {
        unidade: int(valor)
        for valor, unidade in re.findall(r"(\d+)([hms])", tempo)
    }

    return (
        valores.get("h", 0) * 3600 +
        valores.get("m", 0) * 60 +
        valores.get("s", 0)
    )


def segundos_em_tempo(segundos: int) -> str:
    """Formata segundos em 'Xh Ym Zs'."""
    horas = segundos // 3600
    segundos %= 3600
    minutos = segundos // 60
    segundos %= 60

    partes = []
    if horas:
        partes.append(f"{horas}h")
    if minutos:
        partes.append(f"{minutos}m")
    if segundos:
        partes.append(f"{segundos}s")

    return " ".join(partes) if partes else "0s"


def segundos_em_formato_digital(segundos: int) -> str:
    """Formata segundos em 'H:MM:SS'."""
    horas = segundos // 3600
    minutos = (segundos % 3600) // 60
    segundos = segundos % 60
    return f"{horas}:{minutos:02d}:{segundos:02d}"


def segundos_em_formato_hh_mm_ss(segundos: int) -> str:
    """Formata segundos em 'HH:MM:SS' (ex: 3665 -> '01:01:05', 0 -> '00:00:00')."""
    if segundos is None or segundos < 0:
        return "00:00:00"
    horas = segundos // 3600
    minutos = (segundos % 3600) // 60
    segs = segundos % 60
    return f"{horas:02d}:{minutos:02d}:{segs:02d}"


def normalizar_texto(texto: Any) -> str:
    """Normaliza texto removendo acentos, apóstrofos e espaços duplicados em maiúsculas."""
    if not isinstance(texto, str):
        return ""
    texto = texto.strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.replace("'", " ")
    texto = " ".join(texto.split())
    return texto


def normalizar_email(email: Any) -> str:
    """Normaliza e-mail removendo tags <...>, espaços e convertendo para minúsculas."""
    if email is None or pd.isna(email):
        return ""
    txt = str(email).strip().lower()
    m = re.search(r"<([^>]+)>", txt)
    if m:
        return m.group(1).strip().lower()
    return txt


def filtrar_emails_ignorados(df: pd.DataFrame, emails_ignorados: Optional[Collection[str]] = None) -> pd.DataFrame:
    """
    Filtra e remove registros de participantes cujos e-mails constam na lista de ignorados.
    Verifica tanto colunas de e-mail quanto a coluna Nome (caso contenha Nome <email>).
    """
    if not emails_ignorados or df is None or df.empty:
        return df

    emails_set = {normalizar_email(e) for e in emails_ignorados if normalizar_email(e)}
    if not emails_set:
        return df

    col_email = None
    for c in df.columns:
        c_upper = str(c).strip().upper()
        if c_upper in ["EMAIL", "E-MAIL", "CORREIO", "E-MAIL ADDRESS", "EMAIL ADDRESS"]:
            col_email = c
            break
    if col_email is None:
        for c in df.columns:
            c_upper = str(c).strip().upper()
            if "EMAIL" in c_upper or "CORREIO" in c_upper:
                col_email = c
                break

    mascara = pd.Series(True, index=df.index)
    if col_email is not None:
        serie_emails = df[col_email].apply(normalizar_email)
        mascara = mascara & ~serie_emails.isin(emails_set)

    if "Nome" in df.columns:
        serie_nome_email = df["Nome"].apply(normalizar_email)
        mascara = mascara & ~serie_nome_email.isin(emails_set)

    return df[mascara].copy().reset_index(drop=True)


def converter_para_inteiro(val: Any) -> int:
    """Converte valor para inteiro de forma segura, tratando NaN, nulos e strings vazias."""
    if pd.isna(val) or val is None:
        return 0
    if isinstance(val, (int, float)):
        return int(val) if not math.isnan(val) else 0
    val_str = str(val).strip()
    if not val_str:
        return 0
    try:
        return int(float(val_str.replace(",", ".")))
    except (ValueError, TypeError):
        return 0


def extrair_datetime_resumo(texto: Any) -> Optional[datetime]:
    """
    Extrai objeto datetime a partir de strings como '9/14/26,8:01:54 AM', '14/09/2026' ou '2026-09-14'.
    """
    if pd.isna(texto) or not texto:
        return None
    texto_str = str(texto).strip().strip('"\'')
    if not texto_str:
        return None

    formatos = [
        "%m/%d/%y,%I:%M:%S %p",
        "%m/%d/%y, %I:%M:%S %p",
        "%m/%d/%Y,%I:%M:%S %p",
        "%m/%d/%Y, %I:%M:%S %p",
        "%d/%m/%Y,%H:%M:%S",
        "%d/%m/%Y, %H:%M:%S",
        "%d/%m/%y,%H:%M:%S",
        "%d/%m/%y, %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%m/%d/%y",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%Y-%m-%d",
    ]
    for fmt in formatos:
        try:
            return datetime.strptime(texto_str, fmt)
        except ValueError:
            pass

    m = re.match(
        r"^(\d{1,2})/(\d{1,2})/(\d{2,4})(?:[,\s]+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?\s*(AM|PM)?)?",
        texto_str,
        re.IGNORECASE,
    )
    if m:
        p1, p2, p3 = int(m.group(1)), int(m.group(2)), int(m.group(3))
        ano = p3 if p3 > 99 else 2000 + p3
        h = int(m.group(4)) if m.group(4) else 0
        mi = int(m.group(5)) if m.group(5) else 0
        s = int(m.group(6)) if m.group(6) else 0
        ampm = m.group(7)
        if ampm:
            if ampm.upper() == "PM" and h < 12:
                h += 12
            elif ampm.upper() == "AM" and h == 12:
                h = 0

        if p1 > 12:
            dia, mes = p1, p2
        elif p2 > 12:
            mes, dia = p1, p2
        else:
            if ampm:
                mes, dia = p1, p2
            else:
                dia, mes = p1, p2
        try:
            return datetime(ano, mes, dia, h, mi, s)
        except ValueError:
            pass

    return None


def extrair_data_resumo(texto: Any) -> str:
    """
    Extrai e padroniza a data a partir de strings como '9/14/26,8:01:54 AM' para 'DD/MM/AAAA'.
    """
    dt = extrair_datetime_resumo(texto)
    if dt:
        return dt.strftime("%d/%m/%Y")
    if pd.isna(texto) or not texto:
        return ""
    texto_str = str(texto).strip().strip('"\'')
    return texto_str.split(",")[0].strip() if texto_str else ""


def ler_resumo_reuniao(caminho: Path) -> Dict[str, Any]:
    """
    Lê a seção '1. Resumo' da lista de presença do Microsoft Teams.
    Extrai:
    - 'data': data formatada como DD/MM/AAAA
    - 'datetime': objeto datetime da reunião (se disponível)
    - 'duracao_str': duração da reunião (ex: '4h 2m 32s')
    - 'duracao_segundos': duração da reunião em segundos
    - 'tempo_medio_str': tempo médio de participação (ex: '3h 15m 10s')
    - 'tempo_medio_segundos': tempo médio de participação em segundos
    """
    resumo = {
        "data": "",
        "datetime": None,
        "duracao_str": "",
        "duracao_segundos": 0,
        "tempo_medio_str": "",
        "tempo_medio_segundos": 0,
        "tempo_medio_participacao_str": "",
        "tempo_medio_participacao_segundos": 0,
    }
    encodings = ["utf-16", "utf-8-sig", "utf-8", "latin1"]
    linhas_resumo = []

    for enc in encodings:
        try:
            with abrir_arquivo(caminho, mode="r", encoding=enc, errors="replace") as f:
                linhas = f.readlines()
                for linha in linhas:
                    limpa = linha.strip().strip('"\'').lower()
                    if limpa.startswith("2. participantes") or "2. participantes" in limpa:
                        break
                    linhas_resumo.append(linha)
                if linhas_resumo:
                    break
        except Exception:
            continue

    for linha in linhas_resumo:
        linha_str = linha.strip()
        if not linha_str:
            continue

        partes = None
        for sep in ["\t", ";", ","]:
            try:
                reader = csv.reader(io.StringIO(linha_str), delimiter=sep)
                row = next(reader)
                if len(row) >= 2 and row[0].strip():
                    partes = [row[0].strip(), row[1].strip()]
                    break
            except Exception:
                continue

        if not partes:
            for sep in ["\t", ";"]:
                if sep in linha_str:
                    bits = linha_str.split(sep, 1)
                    partes = [bits[0].strip().strip('"\''), bits[1].strip().strip('"\'')]
                    break

        if not partes or len(partes) < 2:
            continue

        chave_norm = normalizar_texto(partes[0])
        valor = partes[1].strip().strip('"\'')

        if "HORA DE INICIO" in chave_norm or "START TIME" in chave_norm:
            if not resumo["data"]:
                dt = extrair_datetime_resumo(valor)
                if dt:
                    resumo["datetime"] = dt
                    resumo["data"] = dt.strftime("%d/%m/%Y")
                else:
                    resumo["data"] = extrair_data_resumo(valor)
        elif "TEMPO MEDIO DE PARTICIPACAO" in chave_norm or "TEMPO MEDIO" in chave_norm or "AVERAGE ATTENDANCE TIME" in chave_norm or "AVERAGE PARTICIPATION TIME" in chave_norm:
            if not resumo["tempo_medio_str"]:
                segs = tempo_em_segundos(valor)
                resumo["tempo_medio_segundos"] = segs
                val_str = valor if valor else (segundos_em_tempo(segs) if segs > 0 else "")
                resumo["tempo_medio_str"] = val_str
                resumo["tempo_medio_participacao_segundos"] = segs
                resumo["tempo_medio_participacao_str"] = val_str
        elif "DURACAO DA REUNIAO" in chave_norm or "DURACAO" in chave_norm or "MEETING DURATION" in chave_norm:
            if not resumo["duracao_str"]:
                segs = tempo_em_segundos(valor)
                resumo["duracao_segundos"] = segs
                resumo["duracao_str"] = valor if valor else (segundos_em_tempo(segs) if segs > 0 else "")

    return resumo


def _chave_natural(texto: str) -> List[Any]:
    """Chave para ordenação natural alfanumérica (ex: 'dia_2' antes de 'dia_10')."""
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r"(\d+)", str(texto))]


def obter_chave_ordenacao_arquivo(caminho: Path) -> Tuple[datetime, List[Any]]:
    """
    Retorna uma tupla (datetime, natural_sort_key) para ordenação cronológica precisa dos encontros.
    Prioridades:
    1. Seção 1 (Resumo da Reunião): 'Hora de início'
    2. Data encontrada no nome do arquivo (YYYY-MM-DD, DD-MM-YYYY, etc.)
    3. Número do dia/encontro no nome do arquivo (ex: dia_1, encontro_2)
    4. Data encontrada dentro das primeiras linhas do arquivo
    5. Fallback com desempate por ordenação natural do nome.
    """
    caminho = Path(caminho)

    # 1. Tenta ler resumo da reunião
    try:
        resumo = ler_resumo_reuniao(caminho)
        if resumo.get("datetime"):
            return (resumo["datetime"], _chave_natural(caminho.name))
        if resumo.get("data"):
            try:
                dt = datetime.strptime(resumo["data"], "%d/%m/%Y")
                return (dt, _chave_natural(caminho.name))
            except Exception:
                pass
    except Exception:
        pass

    nome = caminho.stem

    # 2. Data no nome do arquivo: YYYY-MM-DD ou YYYYMMDD
    m_ano = re.search(r"(\d{4})[-_]?(\d{2})[-_]?(\d{2})", nome)
    if m_ano:
        try:
            ano, mes, dia = int(m_ano.group(1)), int(m_ano.group(2)), int(m_ano.group(3))
            return (datetime(ano, mes, dia), _chave_natural(caminho.name))
        except ValueError:
            pass

    # Data no nome do arquivo: DD-MM-YYYY ou DD_MM_YYYY ou DD.MM.YYYY
    m_dia = re.search(r"(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})", nome)
    if m_dia:
        try:
            dia, mes, ano = int(m_dia.group(1)), int(m_dia.group(2)), int(m_dia.group(3))
            return (datetime(ano, mes, dia), _chave_natural(caminho.name))
        except ValueError:
            pass

    # 3. Número no nome do arquivo (ex: dia_1, encontro_2, aula 3)
    m_num = re.search(r"(?:dia|encontro|aula|reuni[aã]o|modulo|m[oó]dulo)[_\s-]*(\d+)", nome, re.IGNORECASE)
    if m_num:
        num = int(m_num.group(1))
        return (datetime(1970, 1, 1) + timedelta(days=num), _chave_natural(caminho.name))

    # Número isolado no nome (ex: "1.csv", "02 - Presença.csv")
    m_num_iso = re.search(r"^\s*(\d+)", nome)
    if m_num_iso:
        num = int(m_num_iso.group(1))
        return (datetime(1970, 1, 1) + timedelta(days=num), _chave_natural(caminho.name))

    # 4. Tenta encontrar qualquer data nas primeiras 30 linhas do arquivo
    encodings = ["utf-16", "utf-8-sig", "utf-8", "latin1"]
    for enc in encodings:
        try:
            with abrir_arquivo(caminho, mode="r", encoding=enc, errors="replace") as f:
                for _ in range(30):
                    linha = f.readline()
                    if not linha:
                        break
                    dt = extrair_datetime_resumo(linha)
                    if dt:
                        return (dt, _chave_natural(caminho.name))
            break
        except Exception:
            continue

    # 5. Fallback final: datetime máximo com ordenação natural do nome
    return (datetime(9999, 12, 31), _chave_natural(caminho.name))


def ordenar_arquivos_presenca(arquivos: List[Path]) -> List[Path]:
    """Ordena uma lista de arquivos de presença estritamente por data do encontro."""
    if not arquivos:
        return []
    return sorted(arquivos, key=obter_chave_ordenacao_arquivo)



def identificar_colunas_metricas(colunas: List[str]) -> Dict[str, List[str]]:
    """
    Identifica colunas de engajamento e compromisso no Teams:
    - Câmera Ligada
    - Desativar Mudo (Microfone)
    - Levantar as Mãos
    - Reações (Aplauso, Gargalhada, Curtir, Amor, Surpreso, etc.)
    """
    cols_camera = []
    cols_microfone = []
    cols_maos = []
    cols_reacoes = []

    for col in colunas:
        c_norm = normalizar_texto(col)

        # Câmera Ligada
        if "CAMERA LIGADA" in c_norm or ("CAMERA" in c_norm and "COMPROMISSO" in c_norm):
            cols_camera.append(col)
            continue

        # Desativar Mudo (Microfone)
        if "DESATIVAR MUDO" in c_norm or ("MUDO" in c_norm and "COMPROMISSO" in c_norm):
            cols_microfone.append(col)
            continue

        # Levantar as Mãos
        if ("LEVANTAR" in c_norm and "MAO" in c_norm) or ("MAOS" in c_norm and "COMPROMISSO" in c_norm):
            cols_maos.append(col)
            continue

        # Reações
        if (
            "ENVOLVIMENTO: REACAO" in c_norm
            or "REACAO -" in c_norm
            or any(r in c_norm for r in ["APLAUSO", "GARGALHADA", "CURTIR", "AMOR", "SURPRESO"])
            and ("ENVOLVIMENTO" in c_norm or "REACAO" in c_norm)
        ):
            cols_reacoes.append(col)
            continue

    return {
        "camera": cols_camera,
        "microfone": cols_microfone,
        "maos": cols_maos,
        "reacoes": cols_reacoes,
    }


def processar_interacoes_participantes(df_presenca: pd.DataFrame) -> pd.DataFrame:
    """
    Extrai as métricas de interação de cada participante em um único encontro.
    Calcula:
    - camera_qtd: quantidade total de vezes que ligou a câmera
    - camera_ligou: 1 se ligou pelo menos uma vez, 0 caso contrário
    - microfone_qtd: quantidade total de vezes que desativou o mudo
    - microfone_ligou: 1 se desativou mudo pelo menos uma vez, 0 caso contrário
    - maos_qtd: total de vezes que levantou a mão
    - reacoes_qtd: total de reações enviadas
    """
    cols_map = identificar_colunas_metricas(list(df_presenca.columns))

    col_nome = "Nome"
    col_email = "Email" if "Email" in df_presenca.columns else None
    if col_email is None:
        for c in df_presenca.columns:
            if "EMAIL" in str(c).upper() or "CORREIO" in str(c).upper():
                col_email = c
                break

    registros = []
    for _, row in df_presenca.iterrows():
        nome = normalizar_texto(row.get(col_nome, ""))
        if not nome:
            continue
        email = str(row.get(col_email, "")).strip() if col_email else ""
        if email.lower() == "nan":
            email = ""

        cam_val = sum(converter_para_inteiro(row.get(c)) for c in cols_map["camera"])
        mic_val = sum(converter_para_inteiro(row.get(c)) for c in cols_map["microfone"])
        mao_val = sum(converter_para_inteiro(row.get(c)) for c in cols_map["maos"])
        reac_val = sum(converter_para_inteiro(row.get(c)) for c in cols_map["reacoes"])

        registros.append({
            "Nome": nome,
            "Email": email,
            "camera_qtd": cam_val,
            "camera_ligou": 1 if cam_val > 0 else 0,
            "microfone_qtd": mic_val,
            "microfone_ligou": 1 if mic_val > 0 else 0,
            "maos_qtd": mao_val,
            "reacoes_qtd": reac_val,
        })

    df_res = pd.DataFrame(registros)
    if df_res.empty:
        return pd.DataFrame(columns=[
            "Nome", "Email", "camera_qtd", "camera_ligou",
            "microfone_qtd", "microfone_ligou", "maos_qtd", "reacoes_qtd"
        ])

    df_agrupado = df_res.groupby(["Nome", "Email"], as_index=False).agg({
        "camera_qtd": "sum",
        "camera_ligou": lambda s: 1 if (s > 0).any() else 0,
        "microfone_qtd": "sum",
        "microfone_ligou": lambda s: 1 if (s > 0).any() else 0,
        "maos_qtd": "sum",
        "reacoes_qtd": "sum",
    })
    return df_agrupado


def ler_presenca(caminho: Path) -> pd.DataFrame:
    """
    Lê a lista de presença do Teams pulando os cabeçalhos de resumo
    e recortando na seção '3. Atividades em Reunião'.
    Tolerante a diferentes codificações (utf-16, utf-8, latin1) e caminhos longos no Windows.
    """
    encodings = ["utf-16", "utf-8-sig", "utf-8", "latin1"]
    df = None

    for enc in encodings:
        try:
            # Localiza a linha com '2. Participantes' para definir o início dos dados
            skiprows = None
            with abrir_arquivo(caminho, mode="r", encoding=enc, errors="replace") as f:
                linhas_arquivo = f.readlines()
                for idx, linha in enumerate(linhas_arquivo):
                    linha_limpa = linha.strip().strip('"\'').lower()
                    if linha_limpa.startswith("2. participantes") or "2. participantes" in linha_limpa:
                        # Pula até a primeira linha com conteúdo logo abaixo de '2. Participantes'
                        proximo = idx + 1
                        while proximo < len(linhas_arquivo) and not linhas_arquivo[proximo].strip():
                            proximo += 1
                        skiprows = proximo
                        break

            if skiprows is not None:
                for sep in ["\t", ",", ";"]:
                    try:
                        candidato = pd.read_csv(caminho, skiprows=skiprows, encoding=enc, sep=sep)
                        if "Nome" in candidato.columns:
                            df = candidato
                            break
                    except (FileNotFoundError, OSError):
                        try:
                            candidato = pd.read_csv(caminho_estendido(caminho), skiprows=skiprows, encoding=enc, sep=sep)
                            if "Nome" in candidato.columns:
                                df = candidato
                                break
                        except Exception:
                            continue
                    except Exception:
                        continue
                if df is not None and "Nome" in df.columns:
                    break
            else:
                # Fallback: tenta com o padrão fixo de 9 linhas caso não haja o marcador '2. Participantes'
                for sep in ["\t", ",", ";"]:
                    try:
                        candidato = pd.read_csv(caminho, skiprows=9, encoding=enc, sep=sep)
                        if "Nome" in candidato.columns:
                            df = candidato
                            break
                    except (FileNotFoundError, OSError):
                        try:
                            candidato = pd.read_csv(caminho_estendido(caminho), skiprows=9, encoding=enc, sep=sep)
                            if "Nome" in candidato.columns:
                                df = candidato
                                break
                        except Exception:
                            continue
                    except Exception:
                        continue
                if df is not None and "Nome" in df.columns:
                    break
        except Exception:
            continue

    if df is None or "Nome" not in df.columns:
        # Tenta sem skiprows
        for enc in encodings:
            try:
                candidato = pd.read_csv(caminho, encoding=enc, sep=None, engine="python")
                if "Nome" in candidato.columns:
                    df = candidato
                    break
            except (FileNotFoundError, OSError):
                try:
                    candidato = pd.read_csv(caminho_estendido(caminho), encoding=enc, sep=None, engine="python")
                    if "Nome" in candidato.columns:
                        df = candidato
                        break
                except Exception:
                    continue
            except Exception:
                continue

    if df is None or "Nome" not in df.columns:
        raise ValueError(f"Não foi possível identificar a coluna 'Nome' no arquivo de presença: {Path(caminho).name}")

    inicio_segunda_tabela = df[
        df["Nome"].astype(str).str.strip().str.startswith("3. Atividades em Reuni")
    ].index

    if len(inicio_segunda_tabela) > 0:
        df = df.iloc[:inicio_segunda_tabela[0]]

    # Limpa linhas vazias
    df = df.dropna(subset=["Nome"])
    return df


def localizar_coluna_servidor(colunas: Any) -> Any:
    """Localiza a coluna do nome do servidor, priorizando correspondência exata."""
    for col in colunas:
        nome = str(col).strip().upper()
        if nome in ["SERVIDOR", "NOME", "NOME SERVIDOR", "NOME DO SERVIDOR"]:
            return col

    for col in colunas:
        nome = str(col).strip().upper()
        if "SERVIDOR" in nome or "NOME" in nome:
            return col

    return None


def localizar_coluna_uorg(colunas: Any) -> Any:
    """
    Localiza a coluna correspondente a UORG na base de servidores.
    Prioriza correspondência exata para a coluna chamada exatamente 'UORG'
    (ignorando maiúsculas/minúsculas e espaços em branco nas extremidades),
    evitando selecionar colunas auxiliares de código (ex: 'COD UORG' ou 'CÓD UORG').
    """
    # 1. Prioridade absoluta: nome exatamente 'UORG'
    for col in colunas:
        if str(col).strip().upper() == "UORG":
            return col

    # 2. Busca exata por termos equivalentes conhecidos
    for col in colunas:
        nome = str(col).strip().upper()
        if nome in ["LOTACAO", "LOTAÇÃO", "UNIDADE"]:
            return col

    # 3. Busca parcial, descartando explicitamente códigos/IDs (ex: COD UORG, CÓDIGO)
    for col in colunas:
        nome = str(col).strip().upper()
        if any(cod in nome for cod in ["COD", "CÓD", "ID"]):
            continue
        if "UORG" in nome or "LOTACAO" in nome or "LOTAÇÃO" in nome or "UNIDADE" in nome:
            return col

    return None


def juntar_servidor_presenca(servidores: pd.DataFrame, presenca: pd.DataFrame) -> pd.DataFrame:
    """Cruza a presença com os servidores da ANM para obter UORG."""
    servidores = servidores.copy()
    presenca = presenca.copy()

    # Localiza coluna de servidor
    col_servidor = localizar_coluna_servidor(servidores.columns)
    if col_servidor is None:
        col_servidor = servidores.columns[0]

    # Localiza coluna UORG garantindo prioridade para o nome exato "UORG"
    col_uorg = localizar_coluna_uorg(servidores.columns)
    if col_uorg is None:
        col_uorg = servidores.columns[1] if len(servidores.columns) > 1 else col_servidor

    servidores["NOME_NORM"] = servidores[col_servidor].apply(normalizar_texto)
    presenca["Nome"] = presenca["Nome"].apply(normalizar_texto)
    presenca["NOME_NORM"] = presenca["Nome"]

    uorg_map = servidores.drop_duplicates(subset=["NOME_NORM"])[["NOME_NORM", col_uorg]].rename(columns={col_uorg: "UORG"})

    presentes = presenca.merge(
        uorg_map,
        on="NOME_NORM",
        how="left"
    )

    col_duracao = "Duração da Reunião" if "Duração da Reunião" in presentes.columns else None
    if col_duracao is None:
        for c in presentes.columns:
            if "DURA" in str(c).upper() or "PERMANENCIA" in str(c).upper() or "TEMPO" in str(c).upper():
                col_duracao = c
                break
    if col_duracao is None:
        col_duracao = "Duração da Reunião"
        presentes[col_duracao] = "00:00:00"

    col_email = "Email" if "Email" in presentes.columns else None
    if col_email is None:
        for c in presentes.columns:
            if "EMAIL" in str(c).upper() or "CORREIO" in str(c).upper():
                col_email = c
                break
    if col_email is None:
        col_email = "Email"
        presentes[col_email] = ""

    presentes = presentes.rename(columns={col_duracao: "Duração da Reunião", col_email: "Email"})
    return presentes[["Nome", "UORG", "Email", "Duração da Reunião"]]


def ler_convocados(caminho: Path) -> pd.DataFrame:
    """Lê a lista de convocados exportada do Teams (formato 'Nome <email>')."""
    encodings = ["utf-8-sig", "utf-8", "latin1", "utf-16"]
    df = None

    for enc in encodings:
        try:
            with abrir_arquivo(caminho, mode="r", encoding=enc, errors="replace") as f:
                df = pd.read_csv(f, skiprows=2)
            if len(df.columns) > 0:
                break
        except Exception:
            continue

    if df is None:
        for enc in encodings:
            try:
                with abrir_arquivo(caminho, mode="r", encoding=enc, errors="replace") as f:
                    df = pd.read_csv(f)
                if len(df.columns) > 0:
                    break
            except Exception:
                continue

    col_nome = df.columns[0]
    convocados = pd.DataFrame()

    # Extrai Nome e Email
    extraidos = df[col_nome].astype(str).str.extract(r"^(.*?)\s*<([^>]+)>$")
    if extraidos[0].notna().any():
        convocados["Nome"] = extraidos[0].str.strip().apply(normalizar_texto)
        convocados["Email"] = extraidos[1].str.strip()
    else:
        # Se já estiver em colunas separadas
        convocados["Nome"] = df[col_nome].astype(str).str.strip().apply(normalizar_texto)
        if len(df.columns) > 1:
            convocados["Email"] = df[df.columns[1]].astype(str).str.strip()
        else:
            convocados["Email"] = ""

    convocados = convocados.dropna(subset=["Nome"])
    convocados = convocados[convocados["Nome"] != ""]
    return convocados


def marcar_comparecimento(resultado: pd.DataFrame, tempo_comparecimento_minutos: int) -> pd.DataFrame:
    """Marca comparecimento como Presente ou Ausente com base no tempo mínimo."""
    tempo_minimo_segundos = tempo_comparecimento_minutos * 60
    resultado = resultado.copy()

    resultado["Compareceu"] = resultado["Duração da Reunião"].apply(
        lambda tempo: (
            "Presente"
            if (pd.notna(tempo) and tempo_em_segundos(tempo) >= tempo_minimo_segundos)
            else "Ausente"
        )
    )
    return resultado


def criar_dados_do_encontro(resultado: pd.DataFrame, numero_encontro: int) -> pd.DataFrame:
    """Formata colunas do encontro para consolidação."""
    encontro = resultado[["Nome", "Email", "Duração da Reunião", "Compareceu"]].copy()

    encontro["Duração da Reunião"] = encontro["Duração da Reunião"].apply(
        lambda tempo: (
            segundos_em_formato_digital(tempo_em_segundos(tempo))
            if pd.notna(tempo)
            else "00:00:00"
        )
    )

    encontro = encontro.rename(
        columns={
            "Duração da Reunião": f"Encontro {numero_encontro}",
            "Compareceu": f"Comparecimento {numero_encontro}",
        }
    )
    return encontro


# Retrocompatibilidade
criar_dados_do_dia = criar_dados_do_encontro


def calcular_presenca_por_uorg(resultado: pd.DataFrame) -> pd.DataFrame:
    """Calcula quantidade de servidores presentes por UORG no encontro."""
    presentes = resultado[resultado["Compareceu"] == "Presente"]
    if "UORG" in presentes.columns:
        contagem = presentes.groupby("UORG").size().reset_index(name="PRESENTES")
    else:
        contagem = pd.DataFrame(columns=["UORG", "PRESENTES"])
    return contagem


def processar_encontro(
    arquivo: Path,
    servidores: pd.DataFrame,
    numero_encontro: int,
    tempo_comparecimento_minutos: int,
    emails_ignorados: Optional[Collection[str]] = None,
):
    """Processa um arquivo individual de presença de um encontro."""
    dados = ler_presenca(arquivo)
    if emails_ignorados:
        dados = filtrar_emails_ignorados(dados, emails_ignorados)
    resultado = juntar_servidor_presenca(servidores, dados)
    resultado = marcar_comparecimento(resultado, tempo_comparecimento_minutos)
    encontro = criar_dados_do_encontro(resultado, numero_encontro)
    por_uorg = calcular_presenca_por_uorg(resultado)
    return resultado, encontro, por_uorg


# Retrocompatibilidade
processar_dia = processar_encontro


def preparar_convocados(
    caminho_convocados: Path,
    servidores: pd.DataFrame,
    emails_ignorados: Optional[Collection[str]] = None,
) -> pd.DataFrame:
    """Carrega convocados e anexa o UORG a partir da base de servidores."""
    convocados = ler_convocados(caminho_convocados)
    if emails_ignorados:
        convocados = filtrar_emails_ignorados(convocados, emails_ignorados)
    servidores = servidores.copy()

    col_servidor = localizar_coluna_servidor(servidores.columns)
    if col_servidor is None:
        col_servidor = servidores.columns[0]

    col_uorg = localizar_coluna_uorg(servidores.columns)
    if col_uorg is None:
        col_uorg = servidores.columns[1] if len(servidores.columns) > 1 else col_servidor

    servidores["NOME_NORM"] = servidores[col_servidor].apply(normalizar_texto)
    convocados["NOME_NORM"] = convocados["Nome"].apply(normalizar_texto)

    uorg_servidores = servidores[["NOME_NORM", col_uorg]].drop_duplicates(subset=["NOME_NORM"]).rename(columns={col_uorg: "UORG"})
    convocados = convocados.merge(uorg_servidores, on="NOME_NORM", how="left").drop(columns=["NOME_NORM"])
    return convocados


def calcular_total_presencas(resultados: pd.DataFrame) -> pd.DataFrame:
    """Calcula quantidade total de encontros em que o participante esteve 'Presente'."""
    presentes = resultados[resultados["Compareceu"] == "Presente"]
    return (
        presentes[["Nome", "Email"]]
        .value_counts(dropna=False)
        .reset_index(name="Presenças")
    )


def adicionar_total_presencas(convocados: pd.DataFrame, resultados: pd.DataFrame) -> pd.DataFrame:
    """Adiciona a contagem de presenças aos convocados."""
    totais = calcular_total_presencas(resultados)
    relatorio = convocados.merge(totais, on=["Nome", "Email"], how="left")
    relatorio["Presenças"] = relatorio["Presenças"].fillna(0).astype(int)
    return relatorio


def adicionar_encontros_presenca(relatorio: pd.DataFrame, encontros_presenca: List[pd.DataFrame]) -> pd.DataFrame:
    """Adiciona colunas dos encontros (Encontro X e Comparecimento X) ao relatório."""
    relatorio = relatorio.copy()
    for encontro in encontros_presenca:
        # Remove duplicatas se houver
        encontro_limpo = encontro.drop_duplicates(subset=["Nome", "Email"])
        relatorio = relatorio.merge(encontro_limpo, on=["Nome", "Email"], how="left")

    qtd_encontros = len(encontros_presenca)
    for num in range(1, qtd_encontros + 1):
        col_encontro = f"Encontro {num}"
        if col_encontro in relatorio.columns:
            relatorio[col_encontro] = relatorio[col_encontro].fillna("00:00:00")
        elif f"Dia {num}" in relatorio.columns:
            relatorio[f"Dia {num}"] = relatorio[f"Dia {num}"].fillna("00:00:00")

        col_comp = f"Comparecimento {num}"
        if col_comp in relatorio.columns:
            relatorio[col_comp] = relatorio[col_comp].fillna("Ausente")

    return relatorio


# Retrocompatibilidade
adicionar_dias_presenca = adicionar_encontros_presenca


def calcular_carga_participada(linha: pd.Series, quantidade_encontros: int) -> int:
    """Soma tempo total participado em todos os encontros em segundos."""
    total = 0
    for num in range(1, quantidade_encontros + 1):
        tempo = linha.get(f"Encontro {num}")
        if tempo is None:
            tempo = linha.get(f"Dia {num}")
        if pd.notna(tempo):
            total += tempo_em_segundos(tempo)
    return total


def adicionar_carga_horaria(relatorio: pd.DataFrame, quantidade_encontros: int, carga_horaria_horas: float) -> Tuple[pd.DataFrame, pd.Series]:
    """Calcula carga horária participada e percentual em relação à carga total."""
    relatorio = relatorio.copy()
    carga_segundos = relatorio.apply(lambda linha: calcular_carga_participada(linha, quantidade_encontros), axis=1)

    relatorio["CH Total"] = carga_segundos.apply(segundos_em_formato_digital)

    total_segundos_esperado = int(carga_horaria_horas * 3600)
    if total_segundos_esperado > 0:
        relatorio["% da CH Total"] = carga_segundos / total_segundos_esperado
    else:
        relatorio["% da CH Total"] = 0.0

    return relatorio, carga_segundos


def calcular_encontros_minimos(quantidade_encontros: int, percentual_minimo: float) -> int:
    """
    Calcula quantidade mínima de encontros de presença exigidos.
    Usa arredondamento para cima (math.ceil) para garantir que, caso haja
    exigência de presença (percentual > 0), seja necessário pelo menos 1 encontro.
    Evita que cursos de 1 encontro com 75% resultem em 0 encontros mínimos (o que aprovava todos).
    """
    if quantidade_encontros <= 0 or percentual_minimo <= 0:
        return 0
    encontros = math.ceil(round(quantidade_encontros * percentual_minimo, 6))
    return max(1, min(quantidade_encontros, encontros))


# Retrocompatibilidade
calcular_dias_minimos = calcular_encontros_minimos


def adicionar_aptidao(
    relatorio: pd.DataFrame,
    carga_segundos: pd.Series,
    quantidade_encontros: Optional[int] = None,
    carga_horaria_horas: float = 0.0,
    percentual_minimo_carga: float = 0.75,
    percentual_minimo_encontros: float = 0.75,
    quantidade_dias: Optional[int] = None,
    percentual_minimo_dias: Optional[float] = None,
) -> pd.DataFrame:
    """Calcula critérios de aprovação / aptidão."""
    if quantidade_encontros is None:
        quantidade_encontros = quantidade_dias if quantidade_dias is not None else 0
    elif quantidade_dias is not None:
        quantidade_encontros = quantidade_dias

    if percentual_minimo_dias is not None:
        percentual_minimo_encontros = percentual_minimo_dias

    relatorio = relatorio.copy()

    carga_total_seg = int(carga_horaria_horas * 3600)
    carga_minima_seg = int(carga_total_seg * percentual_minimo_carga)
    encontros_minimos = calcular_encontros_minimos(quantidade_encontros, percentual_minimo_encontros)

    status_presenca = (relatorio["Presenças"] >= encontros_minimos).map({True: "APTO", False: "NÃO APTO"})
    relatorio["Apto por CH"] = (carga_segundos >= carga_minima_seg).map({True: "APTO", False: "NÃO APTO"})
    relatorio["Apto por Presença"] = status_presenca
    relatorio["Apto por Encontros"] = status_presenca
    relatorio["Apto por Dias"] = status_presenca
    relatorio["Apto em Ambos"] = (
        (relatorio["Apto por CH"] == "APTO") & (status_presenca == "APTO")
    ).map({True: "APTO", False: "NÃO APTO"})

    return relatorio


def criar_aba_presencas(relatorio: pd.DataFrame, quantidade_encontros: int) -> pd.DataFrame:
    """
    Cria DataFrame para a aba 'Presenças'.
    Contém apenas a coluna 'Nome' e as colunas 'Encontro n' e 'Comparecimento n'.
    """
    colunas_presencas = ["Nome"]
    for num in range(1, quantidade_encontros + 1):
        col_encontro = f"Encontro {num}" if f"Encontro {num}" in relatorio.columns else f"Dia {num}"
        col_comp = f"Comparecimento {num}"
        if col_encontro in relatorio.columns:
            colunas_presencas.append(col_encontro)
        if col_comp in relatorio.columns:
            colunas_presencas.append(col_comp)

    colunas_existentes = [c for c in colunas_presencas if c in relatorio.columns]
    return relatorio[colunas_existentes].copy()


def criar_aba_interacoes_participante(
    relatorio: pd.DataFrame,
    lista_interacoes_encontros: List[pd.DataFrame],
    lista_resultados_encontros: List[pd.DataFrame],
) -> pd.DataFrame:
    """
    Cria a Aba 1: 'Interações por participante'.
    Colunas:
    - Participante
    - Qtd Presenças
    - Câmera
    - % Câmera
    - Microfone
    - % Microfone
    - Mãos levantadas
    - Total Reações
    """
    col_nome = "Nome"
    col_email = "Email" if "Email" in relatorio.columns else None

    # Base de participantes a partir do relatório consolidado
    df_base = relatorio[[col_nome]].copy()
    if col_email:
        df_base["Email"] = relatorio[col_email]

    if "Presenças" in relatorio.columns:
        df_base["Presenças"] = relatorio["Presenças"]
    else:
        df_base["Presenças"] = 0

    dados_participantes = []

    for _, part in df_base.iterrows():
        nome = normalizar_texto(part[col_nome])
        qtd_presencas = int(part["Presenças"])

        presencas_com_camera = 0
        presencas_com_microfone = 0
        total_maos = 0
        total_reacoes = 0

        for num_encontro, (df_int, df_res) in enumerate(
            zip(lista_interacoes_encontros, lista_resultados_encontros), start=1
        ):
            # Verifica comparecimento no encontro
            p_res = df_res[df_res["Nome"] == nome]
            compareceu = False
            if not p_res.empty:
                compareceu = (p_res["Compareceu"] == "Presente").any()

            # Localiza dados de interação do participante no encontro
            p_int = df_int[df_int["Nome"] == nome]
            if not p_int.empty:
                row_int = p_int.iloc[0]
                cam_ligou = int(row_int.get("camera_ligou", 0))
                mic_ligou = int(row_int.get("microfone_ligou", 0))
                maos = int(row_int.get("maos_qtd", 0))
                reac = int(row_int.get("reacoes_qtd", 0))

                total_maos += maos
                total_reacoes += reac

                # Se compareceu ao encontro e ligou a câmera, contabiliza 1 presença com câmera
                if compareceu and cam_ligou >= 1:
                    presencas_com_camera += 1
                elif cam_ligou >= 1 and qtd_presencas > 0:
                    presencas_com_camera += 1

                # Se compareceu ao encontro e desativou mudo, contabiliza 1 presença com microfone
                if compareceu and mic_ligou >= 1:
                    presencas_com_microfone += 1
                elif mic_ligou >= 1 and qtd_presencas > 0:
                    presencas_com_microfone += 1

        if qtd_presencas > 0:
            presencas_com_camera = min(presencas_com_camera, qtd_presencas)
            presencas_com_microfone = min(presencas_com_microfone, qtd_presencas)
            cobertura_cam = presencas_com_camera / qtd_presencas
            cobertura_mic = presencas_com_microfone / qtd_presencas
        else:
            presencas_com_camera = 0
            presencas_com_microfone = 0
            cobertura_cam = 0.0
            cobertura_mic = 0.0

        dados_participantes.append({
            "Participante": nome,
            "Qtd Presenças": qtd_presencas,
            "Câmera": presencas_com_camera,
            "% Câmera": cobertura_cam,
            "Microfone": presencas_com_microfone,
            "% Microfone": cobertura_mic,
            "Mãos levantadas": total_maos,
            "Total Reações": total_reacoes,
        })

    df_aba1 = pd.DataFrame(dados_participantes)
    if "Participante" in df_aba1.columns:
        df_aba1 = df_aba1.sort_values(
            by="Participante",
            ascending=True,
            key=lambda col: col.astype(str).str.strip().str.lower()
        ).reset_index(drop=True)
    return df_aba1


def criar_aba_interacoes_encontro(
    total_convocados: int,
    lista_resumos: List[Dict[str, Any]],
    lista_interacoes_encontros: List[pd.DataFrame],
    lista_resultados_encontros: List[pd.DataFrame],
    config: Optional[Dict[str, Any]] = None,
) -> pd.DataFrame:
    """
    Cria a Aba 2: 'Interações por encontro'.
    Colunas:
    - ID
    - Data
    - CH Realizada
    - Tempo Médio
    - Presentes
    - Ligou a câmera
    - Ligou o microfone
    - Levantou a mão
    - Reações
    """
    linhas = []
    data_inicio_cfg = config.get("DATA_INICIO", "") if config else ""

    for idx, (resumo, df_int, df_res) in enumerate(
        zip(lista_resumos, lista_interacoes_encontros, lista_resultados_encontros),
        start=1
    ):
        data_enc = resumo.get("data") or data_inicio_cfg or ""

        # Duração da reunião em formato digital 00:00:00
        dur_seg = resumo.get("duracao_segundos", 0)
        if dur_seg == 0:
            dur_str_raw = resumo.get("duracao_str", "")
            if dur_str_raw:
                dur_seg = tempo_em_segundos(dur_str_raw)
        if dur_seg == 0 and "Duração da Reunião" in df_res.columns:
            dur_seg = max([tempo_em_segundos(t) for t in df_res["Duração da Reunião"]], default=0)

        dur_str = segundos_em_formato_hh_mm_ss(dur_seg)

        # Tempo médio de participação em formato digital 00:00:00
        tempo_medio_seg = resumo.get("tempo_medio_segundos", 0)
        if tempo_medio_seg == 0:
            tempo_medio_raw = resumo.get("tempo_medio_str", "")
            if tempo_medio_raw:
                tempo_medio_seg = tempo_em_segundos(tempo_medio_raw)
        if tempo_medio_seg == 0 and "Duração da Reunião" in df_res.columns:
            tempos_presentes = [
                tempo_em_segundos(t) for t in df_res["Duração da Reunião"]
                if tempo_em_segundos(t) > 0
            ]
            tempo_medio_seg = int(round(sum(tempos_presentes) / len(tempos_presentes))) if tempos_presentes else 0

        tempo_medio_str = segundos_em_formato_hh_mm_ss(tempo_medio_seg)

        # Presentes em formato 'X de Y' para eliminar aviso de heurística de data do Excel
        presentes_count = int((df_res["Compareceu"] == "Presente").sum()) if "Compareceu" in df_res.columns else 0
        txt_presentes = f"{presentes_count} de {total_convocados}"

        # Ligou a câmera e Ligou o microfone (usuários presentes naquele dia)
        if "Compareceu" in df_res.columns:
            presentes_nomes = set(df_res[df_res["Compareceu"] == "Presente"]["Nome"])
            df_int_presentes = df_int[df_int["Nome"].isin(presentes_nomes)]
            camera_count = int((df_int_presentes["camera_ligou"] >= 1).sum()) if not df_int_presentes.empty else 0
            microfone_count = int((df_int_presentes["microfone_ligou"] >= 1).sum()) if not df_int_presentes.empty else 0
        else:
            camera_count = int((df_int["camera_ligou"] >= 1).sum()) if not df_int.empty else 0
            microfone_count = int((df_int["microfone_ligou"] >= 1).sum()) if not df_int.empty else 0

        txt_camera = f"{camera_count} de {presentes_count}"
        txt_microfone = f"{microfone_count} de {presentes_count}"

        # Mãos levantadas no encontro
        maos_count = int(df_int["maos_qtd"].sum()) if not df_int.empty else 0

        # Reações no encontro
        reacoes_count = int(df_int["reacoes_qtd"].sum()) if not df_int.empty else 0

        linhas.append({
            "ID": idx,
            "Data": data_enc,
            "CH Realizada": dur_str,
            "Tempo Médio": tempo_medio_str,
            "Presentes": txt_presentes,
            "Ligou a câmera": txt_camera,
            "Ligou o microfone": txt_microfone,
            "Levantou a mão": maos_count,
            "Reações": reacoes_count,
        })

    def _parse_data_para_sort(d_str: str) -> datetime:
        if not d_str:
            return datetime(9999, 12, 31)
        try:
            return datetime.strptime(d_str, "%d/%m/%Y")
        except Exception:
            return datetime(9999, 12, 31)

    linhas.sort(key=lambda r: _parse_data_para_sort(r.get("Data", "")))
    for idx_novo, r in enumerate(linhas, start=1):
        r["ID"] = idx_novo

    return pd.DataFrame(linhas)


def criar_aba_indicadores(
    total_convidados: int,
    relatorio: pd.DataFrame,
    df_interacoes_participante: pd.DataFrame,
    df_interacoes_encontro: pd.DataFrame,
    lista_resumos: List[Dict[str, Any]],
    lista_resultados_encontros: List[pd.DataFrame],
) -> pd.DataFrame:
    """
    Cria a Aba 3: 'Indicadores'.
    Colunas: 'Indicador', 'Resultado'
    Linhas:
    - Convidados a participar
    - Participantes totais
    - Participantes certificados
    - Carga horária realizada
    - Tempo médio de participação
    - Taxa de certificação
    - Participações com Câmera
    - Participações com Microfone
    - Total de reações
    - Total de mãos levantadas
    """
    # 1. Convidados a participar
    convidados_val = total_convidados

    # 2. Participantes totais (presentes em pelo menos 1 encontro)
    if "Presenças" in relatorio.columns:
        participantes_totais = int((relatorio["Presenças"] >= 1).sum())
    elif "Qtd Presenças" in df_interacoes_participante.columns:
        participantes_totais = int((df_interacoes_participante["Qtd Presenças"] >= 1).sum())
    elif "Quantidade de presenças" in df_interacoes_participante.columns:
        participantes_totais = int((df_interacoes_participante["Quantidade de presenças"] >= 1).sum())
    else:
        participantes_totais = 0

    # 3. Participantes certificados (por CH)
    if "Apto por CH" in relatorio.columns:
        participantes_certificados = int((relatorio["Apto por CH"] == "APTO").sum())
    elif "Apto ao Certificado" in relatorio.columns:
        participantes_certificados = int((relatorio["Apto ao Certificado"].astype(str).str.upper().isin(["SIM", "APTO"])).sum())
    else:
        participantes_certificados = 0

    # 4. Carga horária realizada (soma duração de reunião de todos os encontros no formato 00:00:00)
    total_segundos_encontros = 0
    total_segundos_tempo_medio = 0
    for idx, (resumo, df_res) in enumerate(zip(lista_resumos, lista_resultados_encontros)):
        seg = resumo.get("duracao_segundos", 0)
        if seg == 0 and "Duração da Reunião" in df_res.columns:
            seg = max([tempo_em_segundos(t) for t in df_res["Duração da Reunião"]], default=0)
        total_segundos_encontros += seg

        seg_tm = resumo.get("tempo_medio_segundos", 0)
        if seg_tm == 0:
            seg_tm_raw = resumo.get("tempo_medio_str", "")
            if seg_tm_raw:
                seg_tm = tempo_em_segundos(seg_tm_raw)
        if seg_tm == 0 and "Duração da Reunião" in df_res.columns:
            tempos = [tempo_em_segundos(t) for t in df_res["Duração da Reunião"] if tempo_em_segundos(t) > 0]
            seg_tm = int(round(sum(tempos) / len(tempos))) if tempos else 0
        total_segundos_tempo_medio += seg_tm

    carga_realizada_str = segundos_em_formato_hh_mm_ss(total_segundos_encontros)
    tempo_medio_realizado_str = segundos_em_formato_hh_mm_ss(total_segundos_tempo_medio)

    # 5. Taxa de certificação (certificados / convidados)
    taxa_certificacao = (
        participantes_certificados / convidados_val
        if convidados_val > 0 else 0.0
    )

    # 6. Participações com Câmera (% de pessoas que ligaram câmera ao menos 1 vez sobre Participantes Totais)
    col_cam = (
        "Câmera" if "Câmera" in df_interacoes_participante.columns
        else ("Presenças com câmera" if "Presenças com câmera" in df_interacoes_participante.columns else None)
    )
    pessoas_camera = int((df_interacoes_participante[col_cam] >= 1).sum()) if col_cam else 0
    participacoes_camera = (
        pessoas_camera / participantes_totais
        if participantes_totais > 0 else 0.0
    )

    # 7. Participações com Microfone (% de pessoas que ligaram microfone ao menos 1 vez sobre Participantes Totais)
    col_mic = (
        "Microfone" if "Microfone" in df_interacoes_participante.columns
        else ("Presenças com microfone" if "Presenças com microfone" in df_interacoes_participante.columns else None)
    )
    pessoas_microfone = int((df_interacoes_participante[col_mic] >= 1).sum()) if col_mic else 0
    participacoes_microfone = (
        pessoas_microfone / participantes_totais
        if participantes_totais > 0 else 0.0
    )

    # 8. Total de reações (soma de todos os encontros e participantes)
    col_reac = (
        "Total Reações" if "Total Reações" in df_interacoes_participante.columns
        else ("Quantidade de reações" if "Quantidade de reações" in df_interacoes_participante.columns else None)
    )
    total_reacoes = int(df_interacoes_participante[col_reac].sum()) if col_reac else 0

    # 9. Total de mãos levantadas (soma de todos os encontros e participantes)
    col_maos = "Mãos levantadas" if "Mãos levantadas" in df_interacoes_participante.columns else None
    total_maos = int(df_interacoes_participante[col_maos].sum()) if col_maos else 0

    linhas = [
        {"Indicador": "Convidados a participar", "Resultado": convidados_val},
        {"Indicador": "Participantes totais", "Resultado": participantes_totais},
        {"Indicador": "Participantes certificados", "Resultado": participantes_certificados},
        {"Indicador": "Carga horária realizada", "Resultado": carga_realizada_str},
        {"Indicador": "Tempo médio de participação", "Resultado": tempo_medio_realizado_str},
        {"Indicador": "Taxa de certificação", "Resultado": taxa_certificacao},
        {"Indicador": "Participações com Câmera", "Resultado": participacoes_camera},
        {"Indicador": "Participações com Microfone", "Resultado": participacoes_microfone},
        {"Indicador": "Total de reações", "Resultado": total_reacoes},
        {"Indicador": "Total de mãos levantadas", "Resultado": total_maos},
    ]

    return pd.DataFrame(linhas)


def criar_aba_custos(
    total_convidados: int,
    relatorio: pd.DataFrame,
    quantidade_encontros: int,
    carga_horaria_horas: float,
    lista_resumos: List[Dict[str, Any]],
    lista_resultados_encontros: List[pd.DataFrame],
    config: Optional[Dict[str, Any]] = None,
) -> pd.DataFrame:
    """
    Cria a Aba 'Custos'.
    Colunas: 'Item', 'Valor'
    Linhas:
    - Convidados a participar
    - Participantes totais
    - Participantes certificados
    - Custo total
    - Quantidade de encontros
    - Carga horária
    - Carga horária realizada
    - Tempo médio de participação realizado
    - Custo por convidado
    - Custo por participante
    - Custo por certificado
    - Custo por encontro
    - Custo por hora prevista
    - Custo por hora realizada
    - Custo por hora de participação média
    """
    convidados_val = int(total_convidados)

    # Participantes totais (presentes em pelo menos 1 encontro)
    if "Presenças" in relatorio.columns:
        participantes_totais = int((relatorio["Presenças"] >= 1).sum())
    elif "Qtd Presenças" in relatorio.columns:
        participantes_totais = int((relatorio["Qtd Presenças"] >= 1).sum())
    elif "Quantidade de presenças" in relatorio.columns:
        participantes_totais = int((relatorio["Quantidade de presenças"] >= 1).sum())
    else:
        participantes_totais = 0

    # Participantes certificados
    if "Apto por CH" in relatorio.columns:
        participantes_certificados = int((relatorio["Apto por CH"] == "APTO").sum())
    elif "Apto ao Certificado" in relatorio.columns:
        participantes_certificados = int((relatorio["Apto ao Certificado"].astype(str).str.upper().isin(["SIM", "APTO"])).sum())
    else:
        participantes_certificados = 0

    # Custo total
    if config and ("CUSTO_TOTAL_EM_REAIS" in config or "CUSTO_TOTAL_REAIS" in config or "CUSTO_TOTAL" in config):
        custo_raw = config.get("CUSTO_TOTAL_EM_REAIS", config.get("CUSTO_TOTAL_REAIS", config.get("CUSTO_TOTAL", 0.0)))
        try:
            if isinstance(custo_raw, str):
                limpo = custo_raw.replace("R$", "").replace("r$", "").strip()
                if "," in limpo and "." in limpo:
                    limpo = limpo.replace(".", "").replace(",", ".")
                elif "," in limpo:
                    limpo = limpo.replace(",", ".")
                custo_total = float(limpo)
            else:
                custo_total = float(custo_raw)
        except (ValueError, TypeError):
            custo_total = 0.0
    else:
        custo_total = 0.0

    qtd_encontros = int(quantidade_encontros)

    # Carga horária prevista (formato HH:MM:SS)
    try:
        ch_segundos = int(float(carga_horaria_horas or 0.0) * 3600)
    except (ValueError, TypeError):
        ch_segundos = 0
    carga_horaria_str = segundos_em_formato_hh_mm_ss(ch_segundos)

    # Carga horária realizada (soma duração real dos encontros no formato HH:MM:SS)
    total_segundos_encontros = 0
    total_segundos_tempo_medio = 0
    for resumo, df_res in zip(lista_resumos, lista_resultados_encontros):
        seg = resumo.get("duracao_segundos", 0)
        if seg == 0 and "Duração da Reunião" in df_res.columns:
            seg = max([tempo_em_segundos(t) for t in df_res["Duração da Reunião"]], default=0)
        total_segundos_encontros += seg

        seg_tm = resumo.get("tempo_medio_segundos", 0)
        if seg_tm == 0:
            seg_tm_raw = resumo.get("tempo_medio_str", "")
            if seg_tm_raw:
                seg_tm = tempo_em_segundos(seg_tm_raw)
        if seg_tm == 0 and "Duração da Reunião" in df_res.columns:
            tempos = [tempo_em_segundos(t) for t in df_res["Duração da Reunião"] if tempo_em_segundos(t) > 0]
            seg_tm = int(round(sum(tempos) / len(tempos))) if tempos else 0
        total_segundos_tempo_medio += seg_tm

    carga_realizada_str = segundos_em_formato_hh_mm_ss(total_segundos_encontros)
    tempo_medio_realizado_str = segundos_em_formato_hh_mm_ss(total_segundos_tempo_medio)

    # Métricas derivadas de custo
    custo_por_convidado = (custo_total / convidados_val) if convidados_val > 0 else 0.0
    custo_por_participante = (custo_total / participantes_totais) if participantes_totais > 0 else 0.0
    custo_por_certificado = (custo_total / participantes_certificados) if participantes_certificados > 0 else 0.0
    custo_por_encontro = (custo_total / qtd_encontros) if qtd_encontros > 0 else 0.0

    # Custo por hora prevista: Custo total / Carga Horária Prevista
    horas_previstas = float(carga_horaria_horas or 0.0)
    if horas_previstas > 0:
        custo_por_hora_prevista = custo_total / horas_previstas
    else:
        custo_por_hora_prevista = 0.0

    # Custo por hora realizada: Custo total / Carga Horária Realizada (com fallback para prevista)
    horas_realizadas = total_segundos_encontros / 3600.0
    if horas_realizadas > 0:
        custo_por_hora_realizada = custo_total / horas_realizadas
    elif horas_previstas > 0:
        custo_por_hora_realizada = custo_por_hora_prevista
    else:
        custo_por_hora_realizada = 0.0

    # Custo por hora de participação média: Custo total / Tempo Médio Acumulado
    horas_tempo_medio = total_segundos_tempo_medio / 3600.0
    if horas_tempo_medio > 0:
        custo_por_hora_participacao = custo_total / horas_tempo_medio
    else:
        custo_por_hora_participacao = 0.0

    linhas = [
        {"Item": "Convidados a participar", "Valor": convidados_val},
        {"Item": "Participantes totais", "Valor": participantes_totais},
        {"Item": "Participantes certificados", "Valor": participantes_certificados},
        {"Item": "Custo total", "Valor": custo_total},
        {"Item": "Quantidade de encontros", "Valor": qtd_encontros},
        {"Item": "Carga horária", "Valor": carga_horaria_str},
        {"Item": "Carga horária realizada", "Valor": carga_realizada_str},
        {"Item": "Tempo médio de participação", "Valor": tempo_medio_realizado_str},
        {"Item": "Custo por convidado", "Valor": custo_por_convidado},
        {"Item": "Custo por participante", "Valor": custo_por_participante},
        {"Item": "Custo por certificado", "Valor": custo_por_certificado},
        {"Item": "Custo por encontro", "Valor": custo_por_encontro},
        {"Item": "Custo por hora prevista", "Valor": custo_por_hora_prevista},
        {"Item": "Custo por hora realizada", "Valor": custo_por_hora_realizada},
        {"Item": "Custo por hora de participação média", "Valor": custo_por_hora_participacao},
    ]

    return pd.DataFrame(linhas)


def formatar_valor_configuracao(chave: str, valor: Any, config: Dict[str, Any]) -> str:
    """Formata valores das variáveis para exibição amigável e não técnica."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        if chave == "DATA_FIM":
            dt_inicio = config.get("DATA_INICIO", "")
            return f"{dt_inicio} (Curso de 1 dia)" if dt_inicio else "Curso de 1 dia"
        if chave == "LOCAL_EMISSAO":
            return "Brasília — DF"
        return "-"

    val_str = str(valor).strip()

    if chave == "CARGA_HORARIA":
        try:
            num = float(val_str.replace(",", "."))
            if num.is_integer():
                return f"{int(num)} horas" if int(num) != 1 else "1 hora"
            return f"{num} horas"
        except (ValueError, TypeError):
            return f"{val_str} horas"

    elif chave in ("PERCENTUAL_MINIMO_CARGA", "PERCENTUAL_MINIMO_ENCONTROS", "PERCENTUAL_MINIMO_DIAS"):
        try:
            pct_num = float(str(val_str).replace("%", "").replace(",", "."))
            if pct_num > 1.0:
                pct_num = pct_num / 100.0
            pct_100 = pct_num * 100.0
            if pct_100.is_integer():
                return f"{int(pct_100)}%"
            return f"{pct_100:.1f}%"
        except (ValueError, TypeError):
            return f"{val_str}%"

    elif chave == "TEMPO_MINIMO_PRESENCA_MINUTOS":
        try:
            minutos = int(float(str(val_str).replace(",", ".")))
            return f"{minutos} minutos" if minutos != 1 else "1 minuto"
        except (ValueError, TypeError):
            return f"{val_str} minutos"

    elif chave == "DATA_FIM":
        if not val_str or val_str == config.get("DATA_INICIO"):
            dt_inicio = config.get("DATA_INICIO", "")
            return f"{dt_inicio} (Curso de 1 dia)" if dt_inicio else "Curso de 1 dia"
        return val_str

    elif chave == "ASSUNTO_EMAIL":
        nome_curso = config.get("NOME_DO_CURSO", "")
        return val_str.replace("{CURSO}", nome_curso).replace("{curso}", nome_curso)

    elif chave in ("CUSTO_TOTAL_EM_REAIS", "CUSTO_TOTAL_REAIS", "CUSTO_TOTAL"):
        try:
            val_limpo = str(val_str).replace("R$", "").replace("r$", "").strip()
            if "," in val_limpo and "." in val_limpo:
                val_limpo = val_limpo.replace(".", "").replace(",", ".")
            elif "," in val_limpo:
                val_limpo = val_limpo.replace(",", ".")
            val_f = float(val_limpo)
            return f"R$ {val_f:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        except (ValueError, TypeError):
            return f"R$ {val_str}"

    return val_str


def criar_aba_informacoes_curso(config: Optional[Dict[str, Any]] = None) -> pd.DataFrame:
    """
    Cria DataFrame para a aba 'Informações do Curso'.
    Contém 3 colunas amigáveis:
    - 'Item': Nome descritivo da informação do curso
    - 'Configuração': Valor formatado de modo acessível e não técnico (ex: '75%', '30 minutos')
    - 'Observação': Breve explicação sobre a finalidade da informação
    """
    if not config:
        config = {}

    pct_encontros = config.get(
        "PERCENTUAL_MINIMO_ENCONTROS",
        config.get("PERCENTUAL_MINIMO_DIAS", 0.75)
    )
    try:
        pct_encontros_num = float(str(pct_encontros).replace("%", "").replace(",", "."))
        if pct_encontros_num > 1.0:
            pct_encontros_num = pct_encontros_num / 100.0
    except (ValueError, TypeError):
        pct_encontros_num = 0.75

    pct_carga = config.get("PERCENTUAL_MINIMO_CARGA", 0.75)
    try:
        pct_carga_num = float(str(pct_carga).replace("%", "").replace(",", "."))
        if pct_carga_num > 1.0:
            pct_carga_num = pct_carga_num / 100.0
    except (ValueError, TypeError):
        pct_carga_num = 0.75

    data_emissao = config.get("DATA_EMISSAO")
    if not data_emissao:
        data_emissao = config.get("DATA_FIM") or config.get("DATA_INICIO") or ""

    linhas = [
        {
            "Item": "Nome do Curso",
            "Configuração": formatar_valor_configuracao("NOME_DO_CURSO", config.get("NOME_DO_CURSO", ""), config),
            "Observação": "Nome oficial da capacitação que será impresso no certificado.",
        },
        {
            "Item": "Carga Horária",
            "Configuração": formatar_valor_configuracao("CARGA_HORARIA", config.get("CARGA_HORARIA", ""), config),
            "Observação": "Carga horária total prevista para o curso, utilizada para o cálculo de frequência.",
        },
        {
            "Item": "Data de Início",
            "Configuração": formatar_valor_configuracao("DATA_INICIO", config.get("DATA_INICIO", ""), config),
            "Observação": "Data do primeiro encontro ou início das atividades do treinamento.",
        },
        {
            "Item": "Data de Fim",
            "Configuração": formatar_valor_configuracao("DATA_FIM", config.get("DATA_FIM", ""), config),
            "Observação": "Data de encerramento das atividades.",
        },
        {
            "Item": "Local de Emissão",
            "Configuração": formatar_valor_configuracao("LOCAL_EMISSAO", config.get("LOCAL_EMISSAO", "Brasília — DF"), config),
            "Observação": "Cidade e UF do órgão emissor exibidas no certificado.",
        },
        {
            "Item": "Data de Emissão",
            "Configuração": formatar_valor_configuracao("DATA_EMISSAO", data_emissao, config),
            "Observação": "Data oficial de registro e expedição impressa no certificado.",
        },
        {
            "Item": "Percentual Mínimo de Carga Horária",
            "Configuração": pct_carga_num,
            "Observação": "Percentual mínimo da carga horária exigido para que o participante seja aprovado.",
        },
        {
            "Item": "Percentual Mínimo de Presença em Encontros",
            "Configuração": pct_encontros_num,
            "Observação": "Percentual mínimo de encontros/aulas em que o participante deve comparecer.",
        },
        {
            "Item": "Tempo Mínimo de Permanência na Reunião",
            "Configuração": formatar_valor_configuracao("TEMPO_MINIMO_PRESENCA_MINUTOS", config.get("TEMPO_MINIMO_PRESENCA_MINUTOS", 30), config),
            "Observação": "Tempo mínimo exigido em cada reunião para validação da presença no encontro.",
        },
    ]

    return pd.DataFrame(linhas)


def organizar_colunas(relatorio: pd.DataFrame, quantidade_encontros: int = 0) -> pd.DataFrame:
    """
    Ordena as colunas da aba principal 'Certificação' de forma limpa e objetiva.
    Remove as colunas 'Encontro n' e 'Comparecimento n' (que agora ficam na aba 'Presenças').
    """
    colunas_principais = ["Nome", "UORG", "Email"]
    colunas_finais = [
        "Presenças",
        "CH Total",
        "% da CH Total",
        "Apto por CH",
        "Apto por Presença" if "Apto por Presença" in relatorio.columns else ("Apto por Encontros" if "Apto por Encontros" in relatorio.columns else "Apto por Dias"),
        "Apto em Ambos",
    ]

    colunas_existentes = [c for c in colunas_principais + colunas_finais if c in relatorio.columns]
    return relatorio[colunas_existentes].copy()


def criar_estilos() -> Dict[str, Any]:
    """Estilos visuais do Excel idênticos aos definidos no notebook."""
    return {
        "cabecalho_fill": PatternFill(fill_type="solid", fgColor="1F4E78"),
        "cabecalho_font": Font(color="FFFFFF", bold=True),
        "linha_par": PatternFill(fill_type="solid", fgColor="F2F2F2"),
        "linha_impar": PatternFill(fill_type="solid", fgColor="FFFFFF"),
        "apto_fill": PatternFill(fill_type="solid", fgColor="C6EFCE"),
        "apto_font": Font(color="006100", bold=True),
        "nao_apto_fill": PatternFill(fill_type="solid", fgColor="FFC7CE"),
        "nao_apto_font": Font(color="9C0006", bold=True),
        "borda_branca": Border(
            left=Side(style="thin", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
            top=Side(style="thin", color="FFFFFF"),
            bottom=Side(style="thin", color="FFFFFF"),
        ),
    }


def alinhar_colunas(planilha):
    """Aplica alinhamentos e formatos de números (como %) nas células."""
    for celula in planilha[1]:
        celula.alignment = Alignment(horizontal="center", vertical="center")

    for num_linha in range(2, planilha.max_row + 1):
        for celula in planilha[num_linha]:
            nome_col = str(planilha.cell(row=1, column=celula.column).value)
            if (
                nome_col in (
                    "Presenças",
                    "Qtd Presenças",
                    "Quantidade de presenças",
                    "Câmera",
                    "Presenças com câmera",
                    "Microfone",
                    "Presenças com microfone",
                    "Mãos levantadas",
                    "Total Reações",
                    "Quantidade de reações",
                    "ID",
                    "Data",
                    "Carga horária",
                    "Carga horária realizada",
                    "CH Realizada",
                    "Tempo médio de participação",
                    "Tempo Médio",
                    "Presentes",
                    "Ligou a câmera",
                    "Ligou o microfone",
                    "Levantou a mão",
                    "Reações",
                    "CH Total",
                )
                or nome_col.startswith("Encontro ")
                or nome_col.startswith("Dia ")
                or nome_col.startswith("Comparecimento")
                or nome_col.startswith("Apto")
            ):
                celula.alignment = Alignment(horizontal="center", vertical="center")
            elif nome_col == "% da CH Total":
                celula.alignment = Alignment(horizontal="center", vertical="center")
                celula.number_format = "0.00%"
            elif nome_col in ("% Câmera", "% Microfone", "Cobertura da câmera", "Cobertura do microfone"):
                celula.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(celula.value, (int, float)):
                    if (celula.value * 100).is_integer():
                        celula.number_format = "0%"
                    else:
                        celula.number_format = "0.0%"
                else:
                    try:
                        v = float(str(celula.value).replace("%", "").strip().replace(",", "."))
                        if v > 1.0:
                            v = v / 100.0
                        celula.value = v
                        celula.number_format = "0.0%"
                    except Exception:
                        pass
            elif nome_col in ("Configuração", "Valor", "Resultado"):
                celula.alignment = Alignment(horizontal="center", vertical="center")
            elif nome_col in ("Observação", "Indicador", "Item"):
                celula.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                celula.alignment = Alignment(horizontal="left", vertical="center")


def formatar_linhas(planilha, estilos):
    """Aplica cores, zebras e bordas."""
    colunas_apto = set()
    colunas_comparecimento = {}

    for celula in planilha[1]:
        nome = str(celula.value)
        if nome.startswith("Apto"):
            colunas_apto.add(celula.column)
        elif nome.startswith("Comparecimento "):
            num_encontro = nome.replace("Comparecimento ", "")
            for outra in planilha[1]:
                if outra.value in (f"Encontro {num_encontro}", f"Dia {num_encontro}"):
                    colunas_comparecimento[celula.column] = outra.column
                    break

    for num_linha, linha in enumerate(planilha.iter_rows(min_row=2), start=2):
        cor_linha = estilos["linha_par"] if num_linha % 2 == 0 else estilos["linha_impar"]
        for celula in linha:
            celula.fill = cor_linha
            if celula.column not in colunas_apto:
                continue

            celula.border = estilos["borda_branca"]
            if celula.value == "APTO":
                celula.fill = estilos["apto_fill"]
                celula.font = estilos["apto_font"]
            elif celula.value == "NÃO APTO":
                celula.fill = estilos["nao_apto_fill"]
                celula.font = estilos["nao_apto_font"]

    for num_linha in range(2, planilha.max_row + 1):
        for col_comp, col_encontro in colunas_comparecimento.items():
            cel_comp = planilha.cell(row=num_linha, column=col_comp)
            cel_encontro = planilha.cell(row=num_linha, column=col_encontro)
            if str(cel_comp.value).strip().upper() == "AUSENTE":
                cel_comp.font = Font(bold=True)
                cel_encontro.font = Font(bold=True)


def formatar_aba_informacoes(planilha, estilos):
    """Aplica formatação específica e refinada para a aba 'Informações do Curso'."""
    borda_celula = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    planilha.row_dimensions[1].height = 28

    for num_linha in range(2, planilha.max_row + 1):
        planilha.row_dimensions[num_linha].height = 24
        cor_fundo = estilos["linha_par"] if num_linha % 2 == 0 else estilos["linha_impar"]

        cel_item = planilha.cell(row=num_linha, column=1)
        cel_conf = planilha.cell(row=num_linha, column=2)
        cel_obs = planilha.cell(row=num_linha, column=3)

        for cel in (cel_item, cel_conf, cel_obs):
            cel.fill = cor_fundo
            cel.border = borda_celula

        cel_item.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        cel_item.alignment = Alignment(horizontal="left", vertical="center")

        item_nome = str(cel_item.value or "")

        # Formata porcentagens como número nativo do Excel com formatação percentual ("0%").
        # Isso elimina completamente o aviso do Excel:
        # "O número nesta célula é formatado como texto ou preenchido por um apóstrofo".
        if "Percentual" in item_nome:
            try:
                val_limpo = str(cel_conf.value or "").replace("%", "").strip().replace(",", ".")
                num_pct = float(val_limpo)
                if num_pct > 1.0:
                    num_pct = num_pct / 100.0
                cel_conf.value = num_pct
                if (num_pct * 100).is_integer():
                    cel_conf.number_format = "0%"
                else:
                    cel_conf.number_format = "0.0%"
            except (ValueError, TypeError):
                pass

        cel_conf.alignment = Alignment(horizontal="center", vertical="center")
        cel_conf.font = Font(name="Calibri", size=11, bold=False, color="000000")

        cel_obs.font = Font(name="Calibri", size=10, italic=False, color="333333")
        cel_obs.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def formatar_aba_indicadores(planilha, estilos):
    """Aplica formatação específica e refinada para a aba 'Indicadores'."""
    borda_celula = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    planilha.row_dimensions[1].height = 28

    for num_linha in range(2, planilha.max_row + 1):
        planilha.row_dimensions[num_linha].height = 24
        cor_fundo = estilos["linha_par"] if num_linha % 2 == 0 else estilos["linha_impar"]

        cel_ind = planilha.cell(row=num_linha, column=1)
        cel_res = planilha.cell(row=num_linha, column=2)

        for cel in (cel_ind, cel_res):
            cel.fill = cor_fundo
            cel.border = borda_celula

        cel_ind.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        cel_ind.alignment = Alignment(horizontal="left", vertical="center")

        ind_nome = str(cel_ind.value or "")

        # Formata porcentagens como número nativo do Excel com formatação percentual ("0%" ou "0.0%")
        if any(termo in ind_nome for termo in ["Taxa", "Participações", "Porcentagem", "%"]):
            try:
                val_limpo = str(cel_res.value or "").replace("%", "").strip().replace(",", ".")
                num_pct = float(val_limpo)
                if num_pct > 1.0:
                    num_pct = num_pct / 100.0
                cel_res.value = num_pct
                if (num_pct * 100).is_integer():
                    cel_res.number_format = "0%"
                else:
                    cel_res.number_format = "0.0%"
            except (ValueError, TypeError):
                pass
            cel_res.alignment = Alignment(horizontal="center", vertical="center")
            cel_res.font = Font(name="Calibri", size=11, bold=True, color="000000")
        else:
            cel_res.alignment = Alignment(horizontal="center", vertical="center")
            cel_res.font = Font(name="Calibri", size=11, bold=False, color="000000")


def formatar_aba_custos(planilha, estilos):
    """Aplica formatação específica e refinada para a aba 'Custos'."""
    borda_celula = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    planilha.row_dimensions[1].height = 28

    for num_linha in range(2, planilha.max_row + 1):
        planilha.row_dimensions[num_linha].height = 24
        cor_fundo = estilos["linha_par"] if num_linha % 2 == 0 else estilos["linha_impar"]

        cel_item = planilha.cell(row=num_linha, column=1)
        cel_val = planilha.cell(row=num_linha, column=2)

        for cel in (cel_item, cel_val):
            cel.fill = cor_fundo
            cel.border = borda_celula

        cel_item.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        cel_item.alignment = Alignment(horizontal="left", vertical="center")

        item_nome = str(cel_item.value or "")

        # Formata custos como moeda brasileira nativa do Excel ("R$" #,##0.00)
        if "Custo" in item_nome:
            try:
                val_limpo = str(cel_val.value or "").replace("R$", "").replace("r$", "").strip().replace(",", ".")
                num_moeda = float(val_limpo)
                cel_val.value = num_moeda
                cel_val.number_format = '"R$" #,##0.00'
            except (ValueError, TypeError):
                pass
            cel_val.alignment = Alignment(horizontal="center", vertical="center")
            cel_val.font = Font(name="Calibri", size=11, bold=(item_nome == "Custo total"), color="000000")
        else:
            cel_val.alignment = Alignment(horizontal="center", vertical="center")
            cel_val.font = Font(name="Calibri", size=11, bold=False, color="000000")


def formatar_aba_presencas(
    planilha,
    workbook,
    estilos: Dict[str, Any],
    datas_encontros: Optional[List[str]] = None,
):
    """
    Aplica formatação específica e refinada para a aba 'Presenças':
    - Mescla os títulos 'Encontro n' e 'Comparecimento n' em um título único por encontro contendo ID e data (ex: 'Encontro 1 — 14/09/2026').
    - Centraliza as colunas de duração e comparecimento.
    - Aplica negrito nas faltas ('Ausente') e tempo zerado ('00:00:00').
    - Define larguras balanceadas para as colunas.
    """
    if datas_encontros is None and "Interações por encontro" in workbook.sheetnames:
        sheet_int = workbook["Interações por encontro"]
        col_data_idx = None
        for col_idx in range(1, sheet_int.max_column + 1):
            if str(sheet_int.cell(row=1, column=col_idx).value).strip() == "Data":
                col_data_idx = col_idx
                break
        if col_data_idx is not None:
            datas_encontros = []
            for r in range(2, sheet_int.max_row + 1):
                val = sheet_int.cell(row=r, column=col_data_idx).value
                datas_encontros.append(str(val).strip() if val is not None else "")

    num_encontros = (planilha.max_column - 1) // 2
    if num_encontros < 1:
        return

    # 1. Mescla os cabeçalhos dos encontros com ID e data
    for idx in range(1, num_encontros + 1):
        col_inicio = 2 * idx
        col_fim = 2 * idx + 1
        if col_fim > planilha.max_column:
            break

        data_str = ""
        if datas_encontros and idx - 1 < len(datas_encontros):
            data_str = str(datas_encontros[idx - 1] or "").strip()

        titulo = f"Encontro {idx} — {data_str}" if data_str else f"Encontro {idx}"

        planilha.cell(row=1, column=col_inicio, value=titulo)
        planilha.cell(row=1, column=col_fim, value=None)

        faixa = f"{get_column_letter(col_inicio)}1:{get_column_letter(col_fim)}1"
        ja_mesclado = any(str(r) == faixa for r in planilha.merged_cells.ranges)
        if not ja_mesclado:
            planilha.merge_cells(start_row=1, start_column=col_inicio, end_row=1, end_column=col_fim)

    # 2. Estilização do cabeçalho
    for col_idx in range(1, planilha.max_column + 1):
        cel = planilha.cell(row=1, column=col_idx)
        cel.fill = estilos["cabecalho_fill"]
        cel.font = estilos["cabecalho_font"]
        cel.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Formatação das linhas de dados (zebra, alinhamento e destaque de 'Ausente')
    for num_linha in range(2, planilha.max_row + 1):
        cor_linha = estilos["linha_par"] if num_linha % 2 == 0 else estilos["linha_impar"]

        for col_idx in range(1, planilha.max_column + 1):
            cel = planilha.cell(row=num_linha, column=col_idx)
            cel.fill = cor_linha
            if col_idx == 1:
                cel.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cel.alignment = Alignment(horizontal="center", vertical="center")

        for idx in range(1, num_encontros + 1):
            col_dur = 2 * idx
            col_comp = 2 * idx + 1
            if col_comp <= planilha.max_column:
                cel_dur = planilha.cell(row=num_linha, column=col_dur)
                cel_comp = planilha.cell(row=num_linha, column=col_comp)
                if str(cel_comp.value).strip().upper() == "AUSENTE":
                    cel_comp.font = Font(name="Calibri", size=11, bold=True, color="000000")
                    cel_dur.font = Font(name="Calibri", size=11, bold=True, color="000000")
                else:
                    cel_comp.font = Font(name="Calibri", size=11, bold=False, color="000000")
                    cel_dur.font = Font(name="Calibri", size=11, bold=False, color="000000")

    # 4. Largura das colunas ajustada proporcionalmente ao tamanho do título mesclado
    maior_nome = max((len(str(c.value)) for c in planilha["A"] if c.value is not None), default=0)
    planilha.column_dimensions["A"].width = min(max(maior_nome + 3, 12), 50)

    for idx in range(1, num_encontros + 1):
        col_dur = 2 * idx
        col_comp = 2 * idx + 1
        if col_comp > planilha.max_column:
            break

        titulo = str(planilha.cell(row=1, column=col_dur).value or "")
        tam_titulo = len(titulo)

        maior_dur = max(
            (len(str(planilha.cell(row=r, column=col_dur).value or "")) for r in range(2, planilha.max_row + 1)),
            default=8,
        )
        maior_comp = max(
            (len(str(planilha.cell(row=r, column=col_comp).value or "")) for r in range(2, planilha.max_row + 1)),
            default=8,
        )

        min_w_dur = max(maior_dur + 3, 10)
        min_w_comp = max(maior_comp + 3, 11)

        w_titulo_total = tam_titulo + 3
        w_total = max(w_titulo_total, min_w_dur + min_w_comp)

        w_dur = max(min_w_dur, math.ceil(w_total / 2))
        w_comp = max(min_w_comp, w_total - w_dur)

        letra_dur = get_column_letter(col_dur)
        letra_comp = get_column_letter(col_comp)

        planilha.column_dimensions[letra_dur].width = w_dur
        planilha.column_dimensions[letra_comp].width = w_comp


def ajustar_largura_colunas(planilha):
    """Ajusta automaticamente a largura de cada coluna."""
    if planilha.title == "Informações do Curso":
        for coluna in planilha.columns:
            letra = coluna[0].column_letter
            maior = 0
            for c in coluna:
                if c.value is None:
                    continue
                # Se for porcentagem numérica (ex: 0.75), calcula o tamanho baseado na representação formatada (ex: '75%')
                if isinstance(c.value, float) and "Percentual" in str(planilha.cell(row=c.row, column=1).value or ""):
                    tam = len(f"{int(c.value * 100)}%")
                else:
                    tam = len(str(c.value))
                if tam > maior:
                    maior = tam

            if letra == "A":
                planilha.column_dimensions[letra].width = max(maior + 4, 38)
            elif letra == "B":
                # Adapta dinamicamente ao maior valor (ex: nome do curso longo), evitando corte de texto
                planilha.column_dimensions[letra].width = max(maior + 6, 32)
            elif letra == "C":
                planilha.column_dimensions[letra].width = min(max(maior + 4, 50), 75)
        return

    if planilha.title in ("Indicadores", "Custos", "Presenças"):
        if planilha.title in ("Indicadores", "Custos"):
            maior = max((len(str(c.value)) for c in planilha["A"] if c.value is not None), default=30)
            planilha.column_dimensions["A"].width = max(maior + 4, 38)
            planilha.column_dimensions["B"].width = 24
        return

    for coluna in planilha.columns:
        maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=0)
        letra = coluna[0].column_letter
        largura = min(maior + 3, 50)
        planilha.column_dimensions[letra].width = max(largura, 12)


def formatar_cabecalho(planilha, estilos):
    """Estiliza linha de cabeçalho."""
    for celula in planilha[1]:
        celula.fill = estilos["cabecalho_fill"]
        celula.font = estilos["cabecalho_font"]


def formatar_workbook(workbook, datas_encontros: Optional[List[str]] = None):
    """Aplica formatação a todas as abas da planilha."""
    estilos = criar_estilos()
    for sheet in workbook.worksheets:
        formatar_cabecalho(sheet, estilos)
        formatar_linhas(sheet, estilos)
        alinhar_colunas(sheet)
        ajustar_largura_colunas(sheet)
        if sheet.title == "Informações do Curso":
            formatar_aba_informacoes(sheet, estilos)
        elif sheet.title == "Indicadores":
            formatar_aba_indicadores(sheet, estilos)
        elif sheet.title == "Custos":
            formatar_aba_custos(sheet, estilos)
        elif sheet.title == "Presenças":
            formatar_aba_presencas(sheet, workbook, estilos, datas_encontros=datas_encontros)


def exportar_excel(
    relatorio: pd.DataFrame,
    caminho_saida: Any,
    df_presencas: Optional[pd.DataFrame] = None,
    df_informacoes: Optional[pd.DataFrame] = None,
    df_interacoes_participante: Optional[pd.DataFrame] = None,
    df_interacoes_encontro: Optional[pd.DataFrame] = None,
    df_indicadores: Optional[pd.DataFrame] = None,
    df_custos: Optional[pd.DataFrame] = None,
    datas_encontros: Optional[List[str]] = None,
):
    """
    Salva planilha formatada contendo:
    - Aba principal 'Certificação'
    - Aba detalhada 'Presenças'
    - Aba 'Interações por participante'
    - Aba 'Interações por encontro'
    - Aba 'Indicadores'
    - Aba 'Custos'
    - Aba informativa 'Informações do Curso'
    """
    if isinstance(caminho_saida, pd.DataFrame) and df_presencas is not None:
        aba_presencas = caminho_saida
        caminho_alvo = Path(df_presencas)
    elif isinstance(caminho_saida, (str, Path)):
        caminho_alvo = Path(caminho_saida)
        aba_presencas = df_presencas
    else:
        caminho_alvo = Path(caminho_saida)
        aba_presencas = None

    if datas_encontros is None and df_interacoes_encontro is not None and "Data" in df_interacoes_encontro.columns:
        datas_encontros = [str(d).strip() if pd.notna(d) else "" for d in df_interacoes_encontro["Data"]]

    nome_seguro = sanitizar_nome_arquivo(caminho_alvo.stem) + caminho_alvo.suffix
    caminho_alvo = caminho_alvo.parent / nome_seguro
    garantir_diretorio(caminho_alvo.parent)

    def _gravar(caminho_final):
        with pd.ExcelWriter(caminho_final, engine="openpyxl") as writer:
            relatorio.to_excel(writer, sheet_name="Certificação", index=False)
            if aba_presencas is not None and not aba_presencas.empty:
                aba_presencas.to_excel(writer, sheet_name="Presenças", index=False)
            if df_interacoes_participante is not None and not df_interacoes_participante.empty:
                df_interacoes_participante.to_excel(writer, sheet_name="Interações por participante", index=False)
            if df_interacoes_encontro is not None and not df_interacoes_encontro.empty:
                df_interacoes_encontro.to_excel(writer, sheet_name="Interações por encontro", index=False)
            if df_indicadores is not None and not df_indicadores.empty:
                df_indicadores.to_excel(writer, sheet_name="Indicadores", index=False)
            if df_custos is not None and not df_custos.empty:
                df_custos.to_excel(writer, sheet_name="Custos", index=False)
            if df_informacoes is not None and not df_informacoes.empty:
                df_informacoes.to_excel(writer, sheet_name="Informações do Curso", index=False)
            formatar_workbook(writer.book, datas_encontros=datas_encontros)

    try:
        _gravar(caminho_alvo)
    except (FileNotFoundError, OSError):
        # Fallback para caminho estendido do Windows (\\?\) caso o caminho ultrapasse MAX_PATH
        caminho_ext = caminho_estendido(caminho_alvo)
        _gravar(caminho_ext)


def gerar_relatorio(
    caminho_servidores: Path,
    caminho_convocados: Path,
    arquivos_presenca: List[Path],
    caminho_saida_excel: Path,
    carga_horaria_horas: float,
    tempo_comparecimento_minutos: int,
    percentual_minimo_carga: float = 0.75,
    percentual_minimo_encontros: float = 0.75,
    percentual_minimo_dias: Optional[float] = None,
    config: Optional[Dict[str, Any]] = None,
) -> pd.DataFrame:
    """Executa o pipeline completo de geração de relatório de presenças."""
    if percentual_minimo_dias is not None:
        percentual_minimo_encontros = percentual_minimo_dias

    caminho_saida_excel = Path(caminho_saida_excel)
    caminho_convocados = Path(caminho_convocados)
    caminho_servidores = Path(caminho_servidores)

    if config is None:
        candidatos_var = []
        if len(caminho_saida_excel.parents) > 2:
            candidatos_var.append(caminho_saida_excel.parents[2] / "variaveis.txt")
        if len(caminho_convocados.parents) > 1:
            candidatos_var.append(caminho_convocados.parents[1] / "variaveis.txt")
        candidatos_var.append(caminho_convocados.parent / "variaveis.txt")
        candidatos_var.append(Path.cwd() / "variaveis.txt")

        for cand in candidatos_var:
            if cand and existe_com_fallback(cand):
                try:
                    from core.config import obter_configuracao
                    config = obter_configuracao(cand.parent)
                    break
                except Exception:
                    try:
                        from core.config import carregar_variaveis
                        config = carregar_variaveis(cand)
                        break
                    except Exception:
                        pass

        if config is None:
            nome_curso_estimado = caminho_saida_excel.stem.replace("Relatório de Presença - ", "").replace("Relatorio", "").strip()
            config = {
                "NOME_DO_CURSO": nome_curso_estimado or "Curso",
                "CARGA_HORARIA": carga_horaria_horas,
                "CUSTO_TOTAL_EM_REAIS": 0.0,
                "DATA_INICIO": "",
                "DATA_FIM": "",
                "LOCAL_EMISSAO": "Brasília — DF",
                "DATA_EMISSAO": "",
                "PERCENTUAL_MINIMO_CARGA": percentual_minimo_carga,
                "PERCENTUAL_MINIMO_ENCONTROS": percentual_minimo_encontros,
                "TEMPO_MINIMO_PRESENCA_MINUTOS": tempo_comparecimento_minutos,
                "IGNORAR_EMAILS": [],
                "EMAILS_IGNORADOS": [],
            }

    emails_ignorados = set()
    if config:
        raw_ign = config.get("IGNORAR_EMAILS") or config.get("EMAILS_IGNORADOS") or []
        if isinstance(raw_ign, str):
            raw_ign = [p.strip() for p in re.split(r"[,;\n]", raw_ign) if p.strip()]
        emails_ignorados = {normalizar_email(e) for e in raw_ign if normalizar_email(e)}
        if emails_ignorados:
            print(f"   E-mails configurados para ignorar: {len(emails_ignorados)}")

    print("-> Carregando lista de servidores...")
    servidores = ler_excel_com_fallback(caminho_servidores)

    arquivos_presenca = ordenar_arquivos_presenca(arquivos_presenca)
    quantidade_encontros = len(arquivos_presenca)
    if quantidade_encontros == 0:
        raise ValueError("Nenhum arquivo de presença foi fornecido.")

    resultados = []
    encontros_presenca = []
    lista_resumos = []
    lista_interacoes = []

    print(f"-> Processando {quantidade_encontros} lista(s) de presença dos encontros...")
    for num_encontro, arq in enumerate(arquivos_presenca, start=1):
        print(f"   [Encontro {num_encontro}] Processando: {arq.name}")
        resultado, encontro, _ = processar_encontro(
            arquivo=arq,
            servidores=servidores,
            numero_encontro=num_encontro,
            tempo_comparecimento_minutos=tempo_comparecimento_minutos,
            emails_ignorados=emails_ignorados,
        )
        resumo = ler_resumo_reuniao(arq)
        dados_brutos = ler_presenca(arq)
        if emails_ignorados:
            dados_brutos = filtrar_emails_ignorados(dados_brutos, emails_ignorados)
        interacoes = processar_interacoes_participantes(dados_brutos)

        resultados.append(resultado)
        encontros_presenca.append(encontro)
        lista_resumos.append(resumo)
        lista_interacoes.append(interacoes)

    todos_resultados = pd.concat(resultados, ignore_index=True)

    print("-> Cruzando dados com a lista de convocados...")
    convocados = preparar_convocados(caminho_convocados, servidores, emails_ignorados=emails_ignorados)

    relatorio = adicionar_total_presencas(convocados, todos_resultados)
    relatorio = adicionar_encontros_presenca(relatorio, encontros_presenca)

    print("-> Calculando carga horária e percentuais...")
    relatorio, carga_segundos = adicionar_carga_horaria(
        relatorio=relatorio,
        quantidade_encontros=quantidade_encontros,
        carga_horaria_horas=carga_horaria_horas,
    )

    print("-> Calculando critérios de aptidão...")
    relatorio = adicionar_aptidao(
        relatorio=relatorio,
        carga_segundos=carga_segundos,
        quantidade_encontros=quantidade_encontros,
        carga_horaria_horas=carga_horaria_horas,
        percentual_minimo_carga=percentual_minimo_carga,
        percentual_minimo_encontros=percentual_minimo_encontros,
    )

    print("-> Formatando e organizando colunas...")
    if "Nome" in relatorio.columns:
        relatorio = relatorio.sort_values(
            by="Nome",
            ascending=True,
            key=lambda col: col.astype(str).str.strip().str.lower(),
        ).reset_index(drop=True)
    elif "Presenças" in relatorio.columns:
        relatorio = relatorio.sort_values(by="Presenças", ascending=False).reset_index(drop=True)

    df_presencas = criar_aba_presencas(relatorio, quantidade_encontros)
    df_certificacao = organizar_colunas(relatorio, quantidade_encontros)
    df_interacoes_participante = criar_aba_interacoes_participante(relatorio, lista_interacoes, resultados)
    df_interacoes_encontro = criar_aba_interacoes_encontro(len(convocados), lista_resumos, lista_interacoes, resultados, config=config)
    df_indicadores = criar_aba_indicadores(len(convocados), relatorio, df_interacoes_participante, df_interacoes_encontro, lista_resumos, resultados)
    df_custos = criar_aba_custos(
        total_convidados=len(convocados),
        relatorio=relatorio,
        quantidade_encontros=quantidade_encontros,
        carga_horaria_horas=carga_horaria_horas,
        lista_resumos=lista_resumos,
        lista_resultados_encontros=resultados,
        config=config,
    )
    df_informacoes = criar_aba_informacoes_curso(config)

    caminho_saida_excel = Path(caminho_saida_excel)
    nome_arquivo_seguro = sanitizar_nome_arquivo(caminho_saida_excel.stem) + caminho_saida_excel.suffix
    caminho_saida_excel = caminho_saida_excel.parent / nome_arquivo_seguro

    datas_encontros = []
    for arq, resumo in zip(arquivos_presenca, lista_resumos):
        dt_res = resumo.get("data")
        if not dt_res:
            try:
                dt, _ = obter_chave_ordenacao_arquivo(arq)
                if dt and dt.year not in (1970, 9999):
                    dt_res = dt.strftime("%d/%m/%Y")
            except Exception:
                pass
        if not dt_res and config and config.get("DATA_INICIO"):
            dt_res = str(config["DATA_INICIO"])
        datas_encontros.append(dt_res or "")

    print(f"-> Formatando e salvando planilha em: {caminho_saida_excel.name}...")
    garantir_diretorio(caminho_saida_excel.parent)
    exportar_excel(
        df_certificacao,
        caminho_saida_excel,
        df_presencas=df_presencas,
        df_informacoes=df_informacoes,
        df_interacoes_participante=df_interacoes_participante,
        df_interacoes_encontro=df_interacoes_encontro,
        df_indicadores=df_indicadores,
        df_custos=df_custos,
        datas_encontros=datas_encontros,
    )

    total_aptos = (df_certificacao["Apto por CH"] == "APTO").sum()
    total_participantes = len(df_certificacao)
    print(f"[OK] Relatório gerado com sucesso!")
    print(f"   Total de participantes convocados: {total_participantes}")
    print(f"   Total de participantes aptos: {total_aptos}")

    return df_certificacao
